import os
import uuid
import asyncio
import io
import calendar
import tempfile
from pathlib import Path
from typing import Optional, List
from urllib.parse import quote, unquote, urlencode
from datetime import datetime, date, timedelta, timezone

from fastapi import APIRouter, Request, Form, Depends, UploadFile, File, BackgroundTasks, HTTPException
from fastapi.responses import RedirectResponse, HTMLResponse, JSONResponse, Response, FileResponse
from app.web.templating import Jinja2Templates
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import case, select, func, or_, delete, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import selectinload
from sqlalchemy.orm.attributes import flag_modified
from aiogram.types import BufferedInputFile
import aiofiles

from app.utils.csrf import generate_csrf_token, validate_csrf, validate_csrf_header
from app.utils.file_manager import delete_file

from app.config import settings
from app.database.core import get_db
from app.database.models import User, Category, Order, OrderItem, Product, CartItem, AuditLog, PromoCode, PromoBanner, ProductImage, SupportChat, SupportMessage, StockNotification, Favorite, PaymeTransaction, ClickTransaction
from app.utils.security import verify_password
from app.bot.loader import bot

from app.database.repositories.users import UserRepository
from app.database.repositories.products import ProductRepository
from app.database.repositories.orders import OrderRepository
from app.services.order_service import OrderRestoreError, OrderService
from app.services.app_settings import (
    find_delivery_option,
    get_card_transfer_details,
    get_delivery_options,
    get_min_order_quantity_for_product,
    get_pickup_address,
    get_pickup_enabled,
    get_shop_all_products_image,
    set_delivery_options,
    set_card_transfer_details,
    set_pickup_address,
    set_pickup_enabled,
    set_shop_all_products_image,
)
from app.services.backup_service import BACKUP_UPLOAD_DIR, BackupBusyError, backup_service
from app.utils.logger import logger
from app.utils.audit import log_action
from app.utils.i18n import tr
from app.utils.product_media import (
    IMAGE_MAX_BYTES,
    IMAGE_EXTENSION,
    delete_product_media_files,
    is_video_file,
    process_product_image,
    public_media_to_path,
    save_uploaded_product_media_batch,
)
from app.utils.upload_security import read_upload_limited, validate_image_upload

router = APIRouter(prefix="/admin", tags=["admin"])
templates = Jinja2Templates(directory="app/templates")


CATEGORY_MEDIA_DIR = Path("media/categories")
CATEGORY_GIF_WIDTH = 432
CATEGORY_GIF_HEIGHT = 510
BANNER_MEDIA_DIR = Path("media/banners")
BANNER_GIF_WIDTH = 640
BANNER_GIF_HEIGHT = 260
CATEGORY_GIF_DURATION_SECONDS = 4
CATEGORY_GIF_FPS = 10
CATEGORY_MEDIA_MAX_BYTES = 50 * 1024 * 1024


def _upload_extension(filename: str) -> str:
    return (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()


def _is_gif_file(upload_file: UploadFile) -> bool:
    return _upload_extension(upload_file.filename or "") == "gif" or upload_file.content_type == "image/gif"


def _category_crop_payload(
    x: Optional[float],
    y: Optional[float],
    width: Optional[float],
    height: Optional[float],
) -> Optional[dict[str, float]]:
    values = (x, y, width, height)
    if any(value is None for value in values):
        return None
    if width <= 0 or height <= 0:
        return None
    return {
        "x": float(x or 0),
        "y": float(y or 0),
        "width": float(width),
        "height": float(height),
    }


async def _probe_media_dimensions(path: Path) -> Optional[tuple[int, int]]:
    cmd = [
        "ffprobe",
        "-v",
        "error",
        "-select_streams",
        "v:0",
        "-show_entries",
        "stream=width,height",
        "-of",
        "csv=s=x:p=0",
        str(path),
    ]
    try:
        proc = await asyncio.create_subprocess_exec(
            *cmd,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.DEVNULL,
        )
        stdout, _ = await proc.communicate()
    except FileNotFoundError:
        logger.warning("ffprobe is not installed; category video/gif conversion skipped")
        return None
    except Exception:
        logger.exception("Could not probe category media")
        return None

    if proc.returncode != 0:
        return None
    raw = stdout.decode("utf-8", errors="ignore").strip().splitlines()
    if not raw or "x" not in raw[0]:
        return None
    try:
        width_raw, height_raw = raw[0].split("x", 1)
        width = int(width_raw)
        height = int(height_raw)
    except (TypeError, ValueError):
        return None
    if width <= 0 or height <= 0:
        return None
    return width, height


def _normalize_category_crop(
    source_width: int,
    source_height: int,
    crop: Optional[dict[str, float]],
    *,
    output_width: int = CATEGORY_GIF_WIDTH,
    output_height: int = CATEGORY_GIF_HEIGHT,
) -> tuple[int, int, int, int]:
    if crop:
        x_limit = max(0, source_width - 1)
        y_limit = max(0, source_height - 1)
        x = max(0, min(x_limit, int(round(crop["x"]))))
        y = max(0, min(y_limit, int(round(crop["y"]))))
        available_width = max(1, source_width - x)
        available_height = max(1, source_height - y)
        width = max(1, min(available_width, int(round(crop["width"]))))
        height = max(1, min(available_height, int(round(crop["height"]))))
        return x, y, width, height

    target_aspect = output_width / output_height
    source_aspect = source_width / source_height
    if source_aspect > target_aspect:
        height = source_height
        width = max(2, int(round(height * target_aspect)))
        x = (source_width - width) // 2
        y = 0
    else:
        width = source_width
        height = max(2, int(round(width / target_aspect)))
        x = 0
        y = (source_height - height) // 2
    return x, y, min(width, source_width - x), min(height, source_height - y)


async def _convert_category_media_to_gif(
    file_bytes: bytes,
    *,
    source_ext: str,
    crop: Optional[dict[str, float]],
    target_dir: Path = CATEGORY_MEDIA_DIR,
    target_prefix: str = "category",
    output_width: int = CATEGORY_GIF_WIDTH,
    output_height: int = CATEGORY_GIF_HEIGHT,
) -> Optional[str]:
    if not file_bytes or len(file_bytes) > CATEGORY_MEDIA_MAX_BYTES:
        return None

    await asyncio.to_thread(target_dir.mkdir, parents=True, exist_ok=True)
    source_ext = source_ext if source_ext in {"gif", "mp4", "webm", "mov", "avi"} else "bin"
    target = target_dir / f"{target_prefix}_{uuid.uuid4()}.gif"

    try:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_root = Path(tmp_dir)
            source = tmp_root / f"source.{source_ext}"
            await asyncio.to_thread(source.write_bytes, file_bytes)

            dimensions = await _probe_media_dimensions(source)
            if not dimensions:
                return None
            crop_x, crop_y, crop_width, crop_height = _normalize_category_crop(
                dimensions[0],
                dimensions[1],
                crop,
                output_width=output_width,
                output_height=output_height,
            )
            filter_complex = (
                f"[0:v]fps={CATEGORY_GIF_FPS},"
                f"crop={crop_width}:{crop_height}:{crop_x}:{crop_y},"
                f"scale={output_width}:{output_height}:flags=lanczos,"
                "split[s0][s1];"
                "[s0]palettegen=stats_mode=single[p];"
                "[s1][p]paletteuse=dither=bayer:bayer_scale=5"
            )
            cmd = [
                "ffmpeg",
                "-y",
                "-t",
                str(CATEGORY_GIF_DURATION_SECONDS),
                "-i",
                str(source),
                "-an",
                "-filter_complex",
                filter_complex,
                "-loop",
                "0",
                str(target),
            ]
            proc = await asyncio.create_subprocess_exec(
                *cmd,
                stdout=asyncio.subprocess.DEVNULL,
                stderr=asyncio.subprocess.DEVNULL,
            )
            rc = await proc.wait()
            if rc != 0 or not target.is_file():
                logger.warning(f"ffmpeg {target_prefix} gif conversion failed with code {rc}")
                return None
    except FileNotFoundError:
        logger.warning("ffmpeg is not installed; category video/gif conversion skipped")
        return None
    except Exception:
        logger.exception("Category media conversion failed")
        return None

    return f"/{target.as_posix()}"


async def _save_category_media(
    image: Optional[UploadFile],
    *,
    crop: Optional[dict[str, float]] = None,
) -> Optional[str]:
    if not image or not image.filename:
        return None

    if is_video_file(image.filename) or _is_gif_file(image):
        try:
            file_bytes = await read_upload_limited(image, CATEGORY_MEDIA_MAX_BYTES)
        except ValueError:
            return None
        return await _convert_category_media_to_gif(
            file_bytes,
            source_ext=_upload_extension(image.filename),
            crop=crop,
        )

    try:
        file_bytes = await read_upload_limited(image, IMAGE_MAX_BYTES)
        processed = await asyncio.to_thread(process_product_image, file_bytes)
    except Exception:
        return None

    unique_name = f"category_{uuid.uuid4()}.{IMAGE_EXTENSION}"
    await asyncio.to_thread(CATEGORY_MEDIA_DIR.mkdir, parents=True, exist_ok=True)
    async with aiofiles.open(CATEGORY_MEDIA_DIR / unique_name, "wb") as f:
        await f.write(processed)
    return f"/{(CATEGORY_MEDIA_DIR / unique_name).as_posix()}"


async def _save_banner_media(
    image: Optional[UploadFile],
    *,
    crop: Optional[dict[str, float]] = None,
) -> Optional[str]:
    if not image or not image.filename:
        return None

    if is_video_file(image.filename) or _is_gif_file(image):
        try:
            file_bytes = await read_upload_limited(image, CATEGORY_MEDIA_MAX_BYTES)
        except ValueError:
            return None
        return await _convert_category_media_to_gif(
            file_bytes,
            source_ext=_upload_extension(image.filename),
            crop=crop,
            target_dir=BANNER_MEDIA_DIR,
            target_prefix="banner",
            output_width=BANNER_GIF_WIDTH,
            output_height=BANNER_GIF_HEIGHT,
        )

    try:
        file_bytes = await read_upload_limited(image, IMAGE_MAX_BYTES)
        processed = await asyncio.to_thread(process_product_image, file_bytes)
    except Exception:
        return None

    unique_name = f"banner_{uuid.uuid4()}.{IMAGE_EXTENSION}"
    await asyncio.to_thread(BANNER_MEDIA_DIR.mkdir, parents=True, exist_ok=True)
    async with aiofiles.open(BANNER_MEDIA_DIR / unique_name, "wb") as f:
        await f.write(processed)
    return f"/{(BANNER_MEDIA_DIR / unique_name).as_posix()}"

def format_datetime_uz(value, format="%d.%m.%Y %H:%M"):
    if value is None:
        return ""
    # Добавляем 5 часов для UTC+5
    local_dt = value + timedelta(hours=5)
    return local_dt.strftime(format)

templates.env.filters["datetime_uz"] = format_datetime_uz

def format_number(value):
    try:
        return f"{int(value):,}".replace(",", " ")
    except (ValueError, TypeError):
        return value

templates.env.filters["number_format"] = format_number

def _support_image_url(path):
    """Convert 'media/support/abc.jpg' to authenticated route '/shop/support/image/abc.jpg'."""
    if not path:
        return path
    fname = path.replace("\\", "/").split("/")[-1]
    return f"/shop/support/image/{fname}"

templates.env.filters["support_image_url"] = _support_image_url
templates.env.globals["hr_workspace_url"] = settings.HR_WORKSPACE_URL or "/hr/"
templates.env.globals["market_workspace_url"] = "/admin"


def _is_safe_login_next(next_url: Optional[str]) -> bool:
    if not next_url:
        return False
    if next_url.startswith("/") and not next_url.startswith("//"):
        return True

    hr_url = (settings.HR_WORKSPACE_URL or "").rstrip("/")
    return bool(
        hr_url
        and hr_url.startswith(("http://", "https://"))
        and (next_url == hr_url or next_url.startswith(f"{hr_url}/"))
    )


def _safe_login_next(next_url: Optional[str]) -> str:
    return next_url if _is_safe_login_next(next_url) else "/admin"


async def _count_unread_new_orders(session: AsyncSession) -> int:
    return (
        await session.execute(
            select(func.count(Order.id)).where(
                Order.status == "new",
                Order.admin_unread == True,
            )
        )
    ).scalar() or 0


def _is_embed_request(request: Request) -> bool:
    return request.query_params.get("embed") == "1"


def _build_admin_url(path: str, *, embed: bool = False, error: Optional[str] = None) -> str:
    params: list[str] = []
    if embed:
        params.append("embed=1")
    if error:
        params.append(f"error={error}")
    if not params:
        return path
    return f"{path}?{'&'.join(params)}"


def _products_popular_return_url(request: Request) -> str:
    params = dict(request.query_params)
    params["modal"] = "popular"
    params.pop("product_id", None)
    return f"{request.url.path}?{urlencode(params)}"


def _safe_products_return_url(return_to: str) -> str:
    if return_to and return_to.startswith("/admin/products") and not return_to.startswith("//"):
        return return_to
    return "/admin/products?modal=popular"


def _format_sum(value: int) -> str:
    return f"{int(value or 0):,}".replace(",", " ") + " сум"


def _product_export_price(product: Product) -> str:
    if product.has_discount:
        return f"{_format_sum(product.discount_price)} (обычная цена: {_format_sum(product.price)})"
    return _format_sum(product.price)


def _product_export_image_bytes(public_path: Optional[str]) -> Optional[io.BytesIO]:
    media_path = public_media_to_path(public_path or "")
    if not media_path or not media_path.is_file():
        return None

    try:
        from PIL import Image as PILImage

        with PILImage.open(media_path) as img:
            img.load()
            if getattr(img, "is_animated", False):
                img.seek(0)
            if img.mode not in ("RGB", "RGBA"):
                img = img.convert("RGBA")
            img.thumbnail((240, 180), PILImage.LANCZOS)
            output = io.BytesIO()
            img.save(output, format="PNG")
            output.seek(0)
            return output
    except Exception:
        logger.exception("Could not prepare product image for XLSX export")
        return None


def _build_client_products_workbook(products: list[Product]) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"

    headers = ["Фото", "Название", "Описание", "Цена"]
    header_fill = PatternFill(start_color="3E2310", end_color="3E2310", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="E0CCB7"),
        right=Side(style="thin", color="E0CCB7"),
        top=Side(style="thin", color="E0CCB7"),
        bottom=Side(style="thin", color="E0CCB7"),
    )

    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 28
    ws.row_dimensions[1].height = 26

    image_refs = []
    for row_idx, product in enumerate(products, start=2):
        ws.row_dimensions[row_idx].height = 92
        description = (product.description_ru or product.description_uz or "").strip()
        values = {
            2: product.name_ru or product.name_uz or f"Товар #{product.id}",
            3: description,
            4: _product_export_price(product),
        }
        for column, value in values.items():
            cell = ws.cell(row=row_idx, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border

        ws.cell(row=row_idx, column=1).border = thin_border
        image_bytes = _product_export_image_bytes(product.image_path)
        if image_bytes:
            xl_image = XLImage(image_bytes)
            image_bytes.seek(0)
            xl_image.width = 112
            xl_image.height = 84
            ws.add_image(xl_image, f"A{row_idx}")
            image_refs.append(image_bytes)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _modal_parent_redirect(url: str) -> HTMLResponse:
    return HTMLResponse(
        f"""<!DOCTYPE html>
<html lang="ru">
<head>
    <meta charset="UTF-8">
    <title>Перенаправление</title>
</head>
<body>
    <script>
        if (window.parent && window.parent !== window) {{
            window.parent.location.href = {url!r};
        }} else {{
            window.location.href = {url!r};
        }}
    </script>
</body>
</html>"""
    )


def _modal_product_saved_response() -> HTMLResponse:
    return HTMLResponse(
        """<!DOCTYPE html>
<html lang="ru">
<head>
    <meta charset="UTF-8">
    <title>Сохранено</title>
</head>
<body>
    <script>
        if (window.parent && window.parent !== window) {
            window.parent.postMessage({ type: "admin-product-saved" }, window.location.origin);
        } else {
            window.location.href = "/admin/products";
        }
    </script>
</body>
</html>"""
    )

async def get_current_admin(request: Request, session: AsyncSession = Depends(get_db)):
    user_id = request.session.get("user_id")
    if not user_id:
        return None
    
    user_repo = UserRepository(session)
    user = await user_repo.get_by_id(user_id)
    
    if user and user.role in ["manager", "superadmin"]:
        if user.role == "manager":
            user.permissions = normalize_permissions(user.permissions)
        return user
    return None

# ── Разделы, к которым можно ограничивать доступ ──
ADMIN_SECTIONS = {
    "dashboard":  "Дашборд",
    "products":   "Товары",
    "categories": "Категории",
    "orders":     "Заказы",
    "users":      "Клиенты",
    "mailing":    "Рассылка",
    "promo":      "Промокоды",
    "banners":    "Баннеры",
    "support":    "Поддержка",
    "audit":      "Журнал",
    "settings":   "Настройки",
    "analytics":  "Аналитика",
    "backup":     "Бэкап",
}

def normalize_permissions(raw_permissions: Optional[dict]) -> dict[str, bool]:
    """Fill missing permission keys and keep backward compatibility for old managers."""
    raw = raw_permissions or {}
    perms = {section_key: bool(raw.get(section_key, False)) for section_key in ADMIN_SECTIONS}

    # Older manager records were created before the dedicated categories section existed.
    # If the key is missing but product access is enabled, keep category access working
    # until the permissions are explicitly re-saved from the admin UI.
    if "categories" not in raw and raw.get("products"):
        perms["categories"] = True

    return perms

def has_permission(user: User, section: str) -> bool:
    """Check if user has permission to access a section. Superadmins always have access."""
    if user.role == "superadmin":
        return True
    perms = normalize_permissions(user.permissions)
    return perms.get(section, False)

def check_permission(request: Request, user: User, section: str):
    """Return error response if user has no access, else None."""
    if has_permission(user, section):
        return None
    return templates.TemplateResponse("admin/error.html", {
        "request": request,
        "message": "Доступ запрещён"
    })


async def _get_assignable_categories(
    session: AsyncSession,
    current_category_id: Optional[int] = None,
) -> list[Category]:
    stmt = select(Category)
    if current_category_id is None:
        stmt = stmt.where(Category.is_active == True)
    else:
        stmt = stmt.where(or_(Category.is_active == True, Category.id == current_category_id))
    stmt = stmt.order_by(Category.is_active.desc(), Category.sort_order.asc(), Category.id.asc())
    return (await session.execute(stmt)).scalars().all()


async def _get_category_for_assignment(
    session: AsyncSession,
    category_id: int,
    *,
    current_category_id: Optional[int] = None,
) -> Optional[Category]:
    category = (await session.execute(select(Category).where(Category.id == category_id))).scalar_one_or_none()
    if not category:
        return None
    if category.is_active or category.id == current_category_id:
        return category
    return None


async def _find_conflicting_category(
    session: AsyncSession,
    *,
    name_ru: str,
    name_uz: str,
    exclude_id: Optional[int] = None,
) -> Optional[Category]:
    normalized_name_ru = name_ru.strip().lower()
    normalized_name_uz = name_uz.strip().lower()
    stmt = select(Category).where(
        or_(
            func.lower(func.trim(Category.name_ru)) == normalized_name_ru,
            func.lower(func.trim(Category.name_uz)) == normalized_name_uz,
        )
    ).limit(1)
    if exclude_id is not None:
        stmt = stmt.where(Category.id != exclude_id)
    return (await session.execute(stmt)).scalar_one_or_none()


def _category_conflict_error(conflict: Optional[Category]) -> str:
    if conflict:
        return f'Категория с таким названием уже существует: «{conflict.name_ru}».'
    return "Категория с таким названием уже существует."

async def _send_stock_notifications(product_id: int, product_name: str):
    """Notify users who subscribed to stock alerts for this product."""
    try:
        from app.database.core import async_session_maker
        async with async_session_maker() as session:
            stmt = (
                select(StockNotification)
                .options(selectinload(StockNotification.user))
                .where(
                    StockNotification.product_id == product_id,
                    StockNotification.notified == False,
                )
            )
            notifications = (await session.execute(stmt)).scalars().all()
            for notif in notifications:
                if notif.user and notif.user.telegram_id:
                    lang = notif.user.language or "ru"
                    text = (
                        f"✅ Товар <b>{product_name}</b> снова в наличии!"
                        if lang == "ru"
                        else f"✅ <b>{product_name}</b> yana mavjud!"
                    )
                    try:
                        await bot.send_message(notif.user.telegram_id, text, parse_mode="HTML")
                    except Exception:
                        pass
                notif.notified = True
            await session.commit()
    except Exception:
        logger.exception("Failed to send stock notifications")

@router.get("/login", response_class=HTMLResponse)
async def login_page(request: Request, next: Optional[str] = None):
    csrf_token = generate_csrf_token(request)
    return templates.TemplateResponse(
        "admin/login.html",
        {
            "request": request,
            "csrf_token": csrf_token,
            "next_url": _safe_login_next(next),
        },
    )

@router.post("/login")
async def login_submit(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    next: str = Form("/admin"),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    user_repo = UserRepository(session)
    user = await user_repo.get_by_login(username)

    if not user or not user.password_hash or not verify_password(password, user.password_hash):
        return templates.TemplateResponse("admin/login.html", {
            "request": request, 
            "error": "Неверный логин или пароль",
            "csrf_token": generate_csrf_token(request),
            "next_url": _safe_login_next(next),
        })
    
    if user.role not in ["manager", "superadmin"]:
         return templates.TemplateResponse("admin/login.html", {
            "request": request, 
            "error": "У вас нет прав доступа",
            "csrf_token": generate_csrf_token(request),
            "next_url": _safe_login_next(next),
        })

    request.session["user_id"] = user.id
    return RedirectResponse(url=_safe_login_next(next), status_code=303)

@router.get("/logout")
async def logout(request: Request):
    request.session.pop("user_id", None)
    return RedirectResponse(url="/admin/login")


@router.get("/api/session")
async def admin_api_session(
    request: Request,
    user: User = Depends(get_current_admin),
):
    headers = {"Cache-Control": "no-store"}
    if not user:
        return JSONResponse(
            {
                "authenticated": False,
                "loginUrl": f"/admin/login?next={quote(settings.HR_WORKSPACE_URL or '/hr/', safe='')}",
                "workspaces": {
                    "market": "/admin",
                    "hr": settings.HR_WORKSPACE_URL or "/hr/",
                },
            },
            status_code=401,
            headers=headers,
        )

    return JSONResponse(
        {
            "authenticated": True,
            "user": {
                "id": user.id,
                "username": user.login or user.username,
                "role": user.role,
                "permissions": normalize_permissions(user.permissions) if user.role == "manager" else {},
            },
            "workspaces": {
                "market": "/admin",
                "hr": settings.HR_WORKSPACE_URL or "/hr/",
            },
        },
        headers=headers,
    )

@router.get("/", response_class=HTMLResponse)
async def dashboard(
    request: Request, 
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse(url="/admin/login")
    denied = check_permission(request, user, "dashboard")
    if denied: return denied

    # 1. KPI: Users Count
    users_count_stmt = select(func.count(User.id)).where(User.role == "user")
    users_count = (await session.execute(users_count_stmt)).scalar() or 0

    # 2. KPI: Orders Today
    uz_now = datetime.now(timezone(timedelta(hours=5)))
    today = uz_now.date()
    today_start_utc = datetime(today.year, today.month, today.day) - timedelta(hours=5)
    today_end_utc = today_start_utc + timedelta(days=1)
    orders_today_stmt = select(func.count(Order.id)).where(
        Order.created_at >= today_start_utc,
        Order.created_at < today_end_utc,
        Order.status != 'cancelled'
    )
    orders_today = (await session.execute(orders_today_stmt)).scalar() or 0

    # 3. KPI: Revenue Month
    current_month = today.month
    current_year = today.year
    revenue_stmt = select(func.sum(Order.total_amount)).where(
        Order.status.in_(['done', 'paid']),
        func.extract('month', Order.created_at) == current_month,
        func.extract('year', Order.created_at) == current_year
    )
    revenue_month = (await session.execute(revenue_stmt)).scalar() or 0

    # 3.1 KPI: Average Order Value
    avg_order_stmt = select(func.avg(Order.total_amount)).where(
        Order.status.in_(['done', 'paid'])
    )
    avg_order_value = (await session.execute(avg_order_stmt)).scalar() or 0

    # 3.2 KPI: Repeat Customers (2+ orders)
    repeat_customers_stmt = select(func.count()).select_from(
        select(Order.user_id)
        .where(Order.status.in_(['done', 'paid']))
        .group_by(Order.user_id)
        .having(func.count(Order.id) >= 2)
        .subquery()
    )
    repeat_customers = (await session.execute(repeat_customers_stmt)).scalar() or 0

    # 4. KPI: Total Debt
    debt_stmt = select(func.sum(User.debt)).where(User.role == "user", User.debt > 0)
    total_debt = (await session.execute(debt_stmt)).scalar() or 0

    # --- CHARTS DATA ---
    
    # Chart 1: Monthly Sales (grouped by year-month)
    sales_stmt = (
        select(
            func.extract('year', Order.created_at).label('year'),
            func.extract('month', Order.created_at).label('month'),
            func.count(Order.id).label('count'),
            func.sum(Order.total_amount).label('sum')
        )
        .where(Order.status.in_(['done', 'paid']))
        .group_by('year', 'month')
        .order_by('year', 'month')
        .limit(12) 
    )
    sales_data_raw = (await session.execute(sales_stmt)).all()
    
    # Format for Chart.js
    monthly_labels = []
    monthly_rev = []
    monthly_count = []
    months_map = {1:'Янв', 2:'Фев', 3:'Мар', 4:'Апр', 5:'Май', 6:'Июн', 7:'Июл', 8:'Авг', 9:'Сен', 10:'Окт', 11:'Ноя', 12:'Дек'}
    
    for row in sales_data_raw:
        m_name = months_map.get(int(row.month), str(row.month))
        monthly_labels.append(f"{m_name} {int(row.year)}")
        monthly_rev.append(int(row.sum))
        monthly_count.append(int(row.count))


    # Chart 2: Top Products
    # Need to join Order and OrderItem (assuming relationship exists or manual join)
    # Let's check models.py later, but assuming OrderItem has product_name and quantity.
    from app.database.models import OrderItem
    top_products_stmt = (
        select(OrderItem.product_name, func.sum(OrderItem.quantity).label('total_qty'))
        .join(Order, OrderItem.order_id == Order.id)
        .where(Order.status.in_(['done', 'paid']))
        .group_by(OrderItem.product_name)
        .order_by(func.sum(OrderItem.quantity).desc())
        .limit(5)
    )
    top_products_raw = (await session.execute(top_products_stmt)).all()
    
    top_prod_labels = [row.product_name for row in top_products_raw]
    top_prod_data = [int(row.total_qty) for row in top_products_raw]

    # Chart 3: Payment Methods
    pay_methods_stmt = (
        select(Order.payment_method, func.count(Order.id))
        .where(Order.status != 'cancelled')
        .group_by(Order.payment_method)
    )
    pay_methods_raw = (await session.execute(pay_methods_stmt)).all()
    
    pay_method_names = {"cash": "Наличные", "card": "Payme", "click": "Click", "card_transfer": "Перевод на карту"}
    pay_labels = [pay_method_names.get(row.payment_method, row.payment_method) for row in pay_methods_raw]
    pay_data = [row[1] for row in pay_methods_raw]

    # Quick insights
    low_stock_stmt = (
        select(Product)
        .where(Product.is_active == True, Product.stock <= 5)
        .order_by(Product.stock.asc())
        .limit(5)
    )
    low_stock_products = (await session.execute(low_stock_stmt)).scalars().all()

    recent_orders_stmt = (
        select(Order)
        .options(selectinload(Order.user))
        .order_by(Order.created_at.desc())
        .limit(6)
    )
    recent_orders = (await session.execute(recent_orders_stmt)).scalars().all()

    status_counts_stmt = (
        select(Order.status, func.count(Order.id))
        .where(Order.status != 'cancelled')
        .group_by(Order.status)
    )
    status_counts_raw = (await session.execute(status_counts_stmt)).all()
    status_counts = {row[0]: row[1] for row in status_counts_raw}

    top_debtors_stmt = (
        select(User)
        .where(User.role == "user", User.debt > 0)
        .order_by(User.debt.desc())
        .limit(5)
    )
    top_debtors = (await session.execute(top_debtors_stmt)).scalars().all()

    top_customers_stmt = (
        select(
            User,
            func.count(Order.id).label("orders_count"),
            func.sum(Order.total_amount).label("total_spent")
        )
        .join(Order, Order.user_id == User.id)
        .where(Order.status.in_(['done', 'paid']))
        .group_by(User.id)
        .order_by(func.sum(Order.total_amount).desc())
        .limit(5)
    )
    top_customers = (await session.execute(top_customers_stmt)).all()

    # Blocked bot users (cached, updated every 24h)
    from app.services.blocked_users import get_blocked_users_count
    blocked_users = await get_blocked_users_count()

    return templates.TemplateResponse("admin/dashboard.html", {
        "request": request,
        "user": user,
        "csrf_token": generate_csrf_token(request),
        
        # KPIs
        "users_count": users_count,
        "orders_today": orders_today,
        "revenue_month": revenue_month,
        "avg_order_value": int(avg_order_value),
        "repeat_customers": repeat_customers,
        "total_debt": total_debt,

        # Charts
        "monthly_labels": monthly_labels,
        "monthly_rev": monthly_rev,
        "monthly_count": monthly_count,
        "top_prod_labels": top_prod_labels,
        "top_prod_data": top_prod_data,
        "pay_labels": pay_labels,
        "pay_data": pay_data,
        "low_stock_products": low_stock_products,
        "recent_orders": recent_orders,
        "status_counts": status_counts,
        "top_debtors": top_debtors,
        "top_customers": top_customers,
        "blocked_users": blocked_users,
    })


def _parse_non_negative_int(value, default: int = 0) -> int:
    try:
        return max(0, int(str(value or "").replace(" ", "")))
    except (TypeError, ValueError):
        return default


@router.get("/settings", response_class=HTMLResponse)
async def admin_settings_page(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "settings")
    if denied:
        return denied

    return templates.TemplateResponse("admin/settings.html", {
        "request": request,
        "user": user,
        "csrf_token": generate_csrf_token(request),
        "pickup_enabled": await get_pickup_enabled(session),
        "pickup_address": await get_pickup_address(session),
        "delivery_options": await get_delivery_options(session),
        "card_transfer_details": await get_card_transfer_details(session),
        "success": request.query_params.get("success"),
        "error": request.query_params.get("error"),
    })


@router.post("/settings")
async def admin_settings_save(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "settings")
    if denied:
        return denied

    form = await request.form()
    pickup_enabled = form.get("pickup_enabled") == "1"
    pickup_address = (form.get("pickup_address") or "").strip()
    card_transfer_number = (form.get("card_transfer_number") or "").strip()
    card_transfer_holder = (form.get("card_transfer_holder") or "").strip()
    delete_ids = set(form.getlist("delete_delivery_option"))
    active_ids = set(form.getlist("delivery_option_active"))

    ids = form.getlist("delivery_option_id[]")
    names_ru = form.getlist("delivery_option_name_ru[]")
    names_uz = form.getlist("delivery_option_name_uz[]")
    prices = form.getlist("delivery_option_price[]")

    options = []
    seen_ids = set()
    for idx, option_id in enumerate(ids):
        option_id = (option_id or "").strip() or f"delivery_{idx + 1}"
        if option_id in delete_ids or option_id in seen_ids:
            continue
        name_ru = (names_ru[idx] if idx < len(names_ru) else "").strip()
        name_uz = (names_uz[idx] if idx < len(names_uz) else "").strip() or name_ru
        if not name_ru:
            continue
        options.append({
            "id": option_id,
            "name_ru": name_ru,
            "name_uz": name_uz,
            "price": _parse_non_negative_int(prices[idx] if idx < len(prices) else 0),
            "is_active": option_id in active_ids,
        })
        seen_ids.add(option_id)

    new_name_ru = (form.get("new_delivery_option_name_ru") or "").strip()
    new_name_uz = (form.get("new_delivery_option_name_uz") or "").strip() or new_name_ru
    if new_name_ru:
        new_id = f"delivery_{uuid.uuid4().hex[:10]}"
        options.append({
            "id": new_id,
            "name_ru": new_name_ru,
            "name_uz": new_name_uz,
            "price": _parse_non_negative_int(form.get("new_delivery_option_price")),
            "is_active": form.get("new_delivery_option_active") == "1",
        })

    if not options:
        return RedirectResponse(
            f"/admin/settings?error={quote('Добавьте хотя бы один вариант доставки.')}",
            status_code=303,
        )
    if not any(option.get("is_active") for option in options):
        return RedirectResponse(
            f"/admin/settings?error={quote('Оставьте хотя бы один активный вариант доставки.')}",
            status_code=303,
        )

    await set_pickup_enabled(session, pickup_enabled)
    await set_pickup_address(session, pickup_address)
    await set_card_transfer_details(
        session,
        card_number=card_transfer_number,
        card_holder=card_transfer_holder,
    )
    await set_delivery_options(session, options)
    await log_action(
        session,
        "Обновление настроек доставки",
        request=request,
        user_id=user.id,
        entity_type="настройки",
        details={
            "самовывоз": pickup_enabled,
            "адрес_самовывоза": pickup_address,
            "вариантов_доставки": len(options),
            "карта_для_перевода": card_transfer_number,
            "владелец_карты": card_transfer_holder,
        },
    )
    await session.commit()

    return RedirectResponse(
        f"/admin/settings?success={quote('Настройки сохранены.')}",
        status_code=303,
    )


@router.get("/categories", response_class=HTMLResponse)
async def categories_list(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "categories")
    if denied:
        return denied

    stats_stmt = (
        select(
            Category,
            func.count(Product.id).label("total_products"),
            func.coalesce(
                func.sum(case((Product.is_active == True, 1), else_=0)),
                0,
            ).label("active_products"),
        )
        .outerjoin(Product, Product.category_id == Category.id)
        .group_by(Category.id)
        .order_by(Category.sort_order.asc(), Category.id.asc())
    )
    rows = (await session.execute(stats_stmt)).all()
    categories = [
        {
            "category": category,
            "total_products": total_products or 0,
            "active_products": active_products or 0,
        }
        for category, total_products, active_products in rows
    ]
    initial_modal = request.query_params.get("modal")
    raw_category_id = request.query_params.get("category_id")
    initial_category_id = int(raw_category_id) if raw_category_id and raw_category_id.isdigit() else None
    all_products_image_path = await get_shop_all_products_image(session)
    all_products_count = (
        await session.execute(
            select(func.count(Product.id))
            .join(Category, Category.id == Product.category_id)
            .where(Product.is_active == True, Category.is_active == True)
        )
    ).scalar() or 0

    return templates.TemplateResponse("admin/categories_list.html", {
        "request": request,
        "user": user,
        "categories": categories,
        "all_products_image_path": all_products_image_path,
        "all_products_count": all_products_count,
        "stats": {
            "total": len(categories),
            "active": sum(1 for row in categories if row["category"].is_active),
            "archived": sum(1 for row in categories if not row["category"].is_active),
        },
        "error": request.query_params.get("error"),
        "success": request.query_params.get("success"),
        "initial_modal": initial_modal,
        "initial_category_id": initial_category_id,
        "csrf_token": generate_csrf_token(request),
    })


@router.post("/categories/all-products-image")
async def all_products_category_image_save(
    request: Request,
    image: UploadFile = File(None),
    remove_image: str = Form(""),
    image_crop_x: Optional[float] = Form(None),
    image_crop_y: Optional[float] = Form(None),
    image_crop_width: Optional[float] = Form(None),
    image_crop_height: Optional[float] = Form(None),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "categories"):
        return RedirectResponse("/admin")

    old_image_path = await get_shop_all_products_image(session)
    media_crop = _category_crop_payload(
        image_crop_x,
        image_crop_y,
        image_crop_width,
        image_crop_height,
    )
    new_image_path = await _save_category_media(image, crop=media_crop)
    if image and image.filename and not new_image_path:
        return RedirectResponse(
            f"/admin/categories?error={quote('Не удалось обработать медиа карточки Все товары. Попробуйте другой файл.')}",
            status_code=303,
        )

    final_image_path = old_image_path
    if new_image_path:
        final_image_path = new_image_path
    elif remove_image == "1":
        final_image_path = None

    await set_shop_all_products_image(session, final_image_path)
    await log_action(
        session,
        "Редактирование карточки Все товары",
        request=request,
        user_id=user.id,
        entity_type="настройка",
        entity_id=0,
        details={"image_path": final_image_path},
    )
    await session.commit()
    if old_image_path and old_image_path != final_image_path:
        await delete_file(old_image_path)

    return RedirectResponse(
        f"/admin/categories?success={quote('Медиа карточки Все товары сохранено.')}",
        status_code=303,
    )


@router.post("/categories/new")
async def category_create(
    request: Request,
    name_ru: str = Form(...),
    name_uz: str = Form(...),
    image: UploadFile = File(None),
    image_crop_x: Optional[float] = Form(None),
    image_crop_y: Optional[float] = Form(None),
    image_crop_width: Optional[float] = Form(None),
    image_crop_height: Optional[float] = Form(None),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "categories"):
        return RedirectResponse("/admin")

    name_ru = name_ru.strip()
    name_uz = name_uz.strip()
    if len(name_ru) < 2 or len(name_uz) < 2:
        return RedirectResponse(
            f"/admin/categories?modal=create&error={quote('Название категории должно быть не короче 2 символов на каждом языке.')}",
            status_code=303,
        )

    conflict = await _find_conflicting_category(session, name_ru=name_ru, name_uz=name_uz)
    if conflict:
        return RedirectResponse(
            f"/admin/categories?modal=create&error={quote(_category_conflict_error(conflict))}",
            status_code=303,
        )

    media_crop = _category_crop_payload(
        image_crop_x,
        image_crop_y,
        image_crop_width,
        image_crop_height,
    )
    image_path = await _save_category_media(image, crop=media_crop)
    if image and image.filename and not image_path:
        return RedirectResponse(
            f"/admin/categories?modal=create&error={quote('Не удалось обработать медиа категории. Попробуйте другой файл.')}",
            status_code=303,
        )

    next_order = (
        await session.execute(select(func.coalesce(func.max(Category.sort_order), 0) + 10))
    ).scalar() or 10
    category = Category(
        name_ru=name_ru,
        name_uz=name_uz,
        image_path=image_path,
        sort_order=next_order,
        is_active=True,
    )
    session.add(category)
    try:
        await session.flush()
        await log_action(
            session,
            "Создание категории",
            request=request,
            user_id=user.id,
            entity_type="категория",
            entity_id=category.id,
            details={"name_ru": name_ru, "name_uz": name_uz},
        )
        await session.commit()
    except IntegrityError:
        await session.rollback()
        if image_path:
            await delete_file(image_path)
        return RedirectResponse(
            f"/admin/categories?modal=create&error={quote('Категория с таким названием уже существует.')}",
            status_code=303,
        )

    return RedirectResponse(
        f"/admin/categories?success={quote('Категория создана.')}",
        status_code=303,
    )


@router.get("/categories/{category_id}/edit", response_class=HTMLResponse)
async def category_edit_form(
    request: Request,
    category_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "categories")
    if denied:
        return denied

    category = (await session.execute(select(Category).where(Category.id == category_id))).scalar_one_or_none()
    if not category:
        return RedirectResponse("/admin/categories", status_code=303)
    return RedirectResponse(
        f"/admin/categories?modal=edit&category_id={category_id}",
        status_code=303,
    )


@router.post("/categories/{category_id}/edit")
async def category_edit_save(
    request: Request,
    category_id: int,
    name_ru: str = Form(...),
    name_uz: str = Form(...),
    image: UploadFile = File(None),
    remove_image: str = Form(""),
    image_crop_x: Optional[float] = Form(None),
    image_crop_y: Optional[float] = Form(None),
    image_crop_width: Optional[float] = Form(None),
    image_crop_height: Optional[float] = Form(None),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "categories"):
        return RedirectResponse("/admin")

    category = (
        await session.execute(
            select(Category).where(Category.id == category_id).with_for_update()
        )
    ).scalar_one_or_none()
    if not category:
        return RedirectResponse("/admin/categories", status_code=303)

    name_ru = name_ru.strip()
    name_uz = name_uz.strip()
    if len(name_ru) < 2 or len(name_uz) < 2:
        return RedirectResponse(
            f"/admin/categories?modal=edit&category_id={category_id}&error={quote('Название категории должно быть не короче 2 символов на каждом языке.')}",
            status_code=303,
        )

    conflict = await _find_conflicting_category(
        session,
        name_ru=name_ru,
        name_uz=name_uz,
        exclude_id=category_id,
    )
    if conflict:
        return RedirectResponse(
            f"/admin/categories?modal=edit&category_id={category_id}&error={quote(_category_conflict_error(conflict))}",
            status_code=303,
        )

    category.name_ru = name_ru
    category.name_uz = name_uz
    old_image_path = category.image_path
    media_crop = _category_crop_payload(
        image_crop_x,
        image_crop_y,
        image_crop_width,
        image_crop_height,
    )
    new_image_path = await _save_category_media(image, crop=media_crop)
    if image and image.filename and not new_image_path:
        return RedirectResponse(
            f"/admin/categories?modal=edit&category_id={category_id}&error={quote('Не удалось обработать медиа категории. Попробуйте другой файл.')}",
            status_code=303,
        )
    if new_image_path:
        category.image_path = new_image_path
    elif remove_image == "1":
        category.image_path = None

    try:
        await log_action(
            session,
            "Редактирование категории",
            request=request,
            user_id=user.id,
            entity_type="категория",
            entity_id=category.id,
            details={"name_ru": name_ru, "name_uz": name_uz},
        )
        await session.commit()
    except IntegrityError:
        await session.rollback()
        if new_image_path:
            await delete_file(new_image_path)
        return RedirectResponse(
            f"/admin/categories?modal=edit&category_id={category_id}&error={quote('Категория с таким названием уже существует.')}",
            status_code=303,
        )
    if old_image_path and old_image_path != category.image_path:
        await delete_file(old_image_path)

    return RedirectResponse(
        f"/admin/categories?success={quote('Изменения сохранены.')}",
        status_code=303,
    )


@router.post("/categories/reorder", response_class=JSONResponse, dependencies=[Depends(validate_csrf_header)])
async def categories_reorder(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return JSONResponse({"error": "unauthorized"}, status_code=401)
    if not has_permission(user, "categories"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    try:
        payload = await request.json()
        category_ids = [int(item) for item in payload.get("ids", [])]
    except (TypeError, ValueError, AttributeError):
        return JSONResponse({"error": "invalid_payload"}, status_code=400)

    if not category_ids:
        return JSONResponse({"error": "empty_order"}, status_code=400)

    categories = (
        await session.execute(
            select(Category)
            .where(Category.id.in_(category_ids))
            .with_for_update()
        )
    ).scalars().all()
    category_by_id = {category.id: category for category in categories}
    if len(category_by_id) != len(set(category_ids)):
        return JSONResponse({"error": "category_not_found"}, status_code=404)

    for index, category_id in enumerate(category_ids, start=1):
        category_by_id[category_id].sort_order = index * 10

    await session.commit()
    return JSONResponse({"ok": True})


@router.post("/categories/{category_id}/toggle")
async def category_toggle_status(
    request: Request,
    category_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "categories"):
        return RedirectResponse("/admin")

    category = (
        await session.execute(
            select(Category).where(Category.id == category_id).with_for_update()
        )
    ).scalar_one_or_none()
    if not category:
        return RedirectResponse("/admin/categories", status_code=303)

    total_products = (
        await session.execute(
            select(func.count(Product.id)).where(Product.category_id == category_id)
        )
    ).scalar() or 0

    if category.is_active and total_products > 0:
        return RedirectResponse(
            f"/admin/categories?error={quote('Нельзя архивировать категорию, пока в ней есть товары. Сначала перенесите их в другую категорию.')}",
            status_code=303,
        )

    category.is_active = not category.is_active
    await log_action(
        session,
        "Смена статуса категории",
        request=request,
        user_id=user.id,
        entity_type="категория",
        entity_id=category.id,
        details={"is_active": category.is_active},
    )
    await session.commit()

    success_message = "Категория восстановлена." if category.is_active else "Категория архивирована."
    return RedirectResponse(
        f"/admin/categories?success={quote(success_message)}",
        status_code=303,
    )


@router.post("/categories/{category_id}/delete")
async def category_delete(
    request: Request,
    category_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "categories"):
        return RedirectResponse("/admin")

    category = (
        await session.execute(
            select(Category).where(Category.id == category_id).with_for_update()
        )
    ).scalar_one_or_none()
    if not category:
        return RedirectResponse("/admin/categories", status_code=303)

    total_products = (
        await session.execute(
            select(func.count(Product.id)).where(Product.category_id == category_id)
        )
    ).scalar() or 0
    if total_products > 0:
        return RedirectResponse(
            f"/admin/categories?error={quote('Удалить можно только пустую категорию без товаров.')}",
            status_code=303,
        )

    await log_action(
        session,
        "Удаление категории",
        request=request,
        user_id=user.id,
        entity_type="категория",
        entity_id=category.id,
        details={"name_ru": category.name_ru, "name_uz": category.name_uz},
    )
    image_path = category.image_path
    await session.delete(category)
    await session.commit()
    if image_path:
        await delete_file(image_path)

    return RedirectResponse(
        f"/admin/categories?success={quote('Категория удалена.')}",
        status_code=303,
    )

@router.get("/products", response_class=HTMLResponse)
async def products_list(
    request: Request,
    q: str = "",
    status: str = "active",
    stock: str = "all",
    category_id: str = "",
    modal: str = "",
    product_id: str = "",
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "products")
    if denied: return denied

    error = request.query_params.get("error")
    if error == "invalid_stock":
        error = "Остаток не может быть отрицательным."
    elif error == "empty_client_export":
        error = "Выберите хотя бы один товар для Excel-подборки."

    selected_category_id = int(category_id) if category_id and category_id.isdigit() else None
    initial_modal = modal if modal in {"create", "edit", "popular"} else ""
    initial_product_id = int(product_id) if product_id and product_id.isdigit() else None

    stmt = select(Product)
    if selected_category_id:
        stmt = stmt.where(Product.category_id == selected_category_id)

    if q:
        safe_query = q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        stmt = stmt.where(
            or_(
                Product.name_ru.ilike(f"%{safe_query}%", escape="\\"),
                Product.name_uz.ilike(f"%{safe_query}%", escape="\\")
            )
        )

    if status == "active":
        stmt = stmt.where(Product.is_active == True)
    elif status == "inactive":
        stmt = stmt.where(Product.is_active == False)

    if stock == "low":
        stmt = stmt.where(Product.stock <= 5)
    elif stock == "out":
        stmt = stmt.where(Product.stock <= 0)

    stmt = stmt.order_by(Product.id.desc())
    products = (await session.execute(stmt)).scalars().all()
    categories = (
        await session.execute(
            select(Category).order_by(Category.is_active.desc(), Category.sort_order.asc(), Category.id.asc())
        )
    ).scalars().all()

    total_products = (await session.execute(select(func.count(Product.id)))).scalar() or 0
    active_count = (await session.execute(select(func.count(Product.id)).where(Product.is_active == True))).scalar() or 0
    inactive_count = (await session.execute(select(func.count(Product.id)).where(Product.is_active == False))).scalar() or 0
    popular_count = (await session.execute(select(func.count(Product.id)).where(Product.is_popular == True))).scalar() or 0
    low_stock_count = (await session.execute(
        select(func.count(Product.id)).where(Product.is_active == True, Product.stock <= 5)
    )).scalar() or 0
    popular_products = (
        await session.execute(
            select(Product)
            .options(selectinload(Product.category))
            .where(Product.is_popular == True)
            .order_by(Product.popular_sort_order.asc(), Product.id.desc())
        )
    ).scalars().all()
    popular_ids = {product.id for product in popular_products}

    popular_catalog_stmt = (
        select(Product)
        .options(selectinload(Product.category))
        .join(Category, Category.id == Product.category_id)
    )
    visible_condition = (Product.is_active == True) & (Category.is_active == True)
    if popular_ids:
        popular_catalog_stmt = popular_catalog_stmt.where(
            or_(visible_condition, Product.id.in_(popular_ids))
        )
    else:
        popular_catalog_stmt = popular_catalog_stmt.where(visible_condition)
    popular_catalog_stmt = popular_catalog_stmt.order_by(
        case((Product.is_popular == True, 0), else_=1),
        Product.name_ru.asc(),
    )
    popular_catalog_products = (await session.execute(popular_catalog_stmt)).scalars().all()
    client_catalog_products = (
        await session.execute(
            select(Product)
            .options(selectinload(Product.category))
            .join(Category, Category.id == Product.category_id)
            .where(Product.is_active == True, Category.is_active == True)
            .order_by(Product.name_ru.asc(), Product.id.asc())
        )
    ).scalars().all()
    return templates.TemplateResponse("admin/products_list.html", {
        "request": request,
        "user": user,
        "products": products,
        "categories": categories,
        "popular_products": popular_products,
        "popular_product_ids": popular_ids,
        "popular_catalog_products": popular_catalog_products,
        "client_catalog_products": client_catalog_products,
        "popular_return_url": _products_popular_return_url(request),
        "filters": {"q": q, "status": status, "stock": stock, "category_id": selected_category_id},
        "stats": {
            "total": total_products,
            "active": active_count,
            "inactive": inactive_count,
            "popular": popular_count,
            "low_stock": low_stock_count
        },
        "error": error,
        "csrf_token": generate_csrf_token(request),
        "initial_modal": initial_modal,
        "initial_product_id": initial_product_id,
    })


@router.post("/products/client-export")
async def products_client_export(
    request: Request,
    product_ids: List[int] = Form(default=[]),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "products"):
        return RedirectResponse("/admin")

    selected_ids = []
    seen_ids = set()
    for product_id in product_ids or []:
        if product_id not in seen_ids:
            selected_ids.append(product_id)
            seen_ids.add(product_id)

    if not selected_ids:
        return RedirectResponse("/admin/products?error=empty_client_export", status_code=303)

    products = (
        await session.execute(
            select(Product)
            .options(selectinload(Product.category))
            .join(Category, Category.id == Product.category_id)
            .where(
                Product.id.in_(selected_ids),
                Product.is_active == True,
                Category.is_active == True,
            )
        )
    ).scalars().all()
    product_by_id = {product.id: product for product in products}
    ordered_products = [product_by_id[product_id] for product_id in selected_ids if product_id in product_by_id]
    if not ordered_products:
        return RedirectResponse("/admin/products?error=empty_client_export", status_code=303)

    output = await asyncio.to_thread(_build_client_products_workbook, ordered_products)
    await log_action(
        session,
        "Экспорт подборки товаров для клиента",
        request=request,
        user_id=user.id,
        entity_type="товар",
        details={"product_ids": selected_ids},
    )
    await session.commit()

    from starlette.responses import StreamingResponse

    filename = f"client_products_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

@router.get("/products/new", response_class=HTMLResponse)
async def product_create_form(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "products")
    if denied: return denied

    embed_mode = _is_embed_request(request)
    categories = await _get_assignable_categories(session)
    if not categories:
        if has_permission(user, "categories"):
            target_url = f"/admin/categories?error={quote('Сначала создайте хотя бы одну активную категорию.')}"
        else:
            target_url = f"/admin/products?error={quote('Нельзя создать товар, пока нет активных категорий. Обратитесь к администратору с доступом к разделу категорий.')}"
        if embed_mode:
            return _modal_parent_redirect(target_url)
        return RedirectResponse(
            target_url,
            status_code=303,
        )

    if not embed_mode:
        return RedirectResponse("/admin/products?modal=create", status_code=303)

    return templates.TemplateResponse("admin/product_edit.html", {
        "request": request,
        "user": user,
        "categories": categories,
        "product": None,
        "extra_images": [],
        "csrf_token": generate_csrf_token(request),
        "embed": embed_mode,
        "layout_template": "admin/embed_base.html" if embed_mode else "admin/base.html",
    })

from app.web.schemas.products import ProductCreateSchema
from fastapi import UploadFile, File

def _has_invalid_product_discount(product_data: ProductCreateSchema) -> bool:
    return bool(
        product_data.discount_price is not None
        and product_data.discount_price >= product_data.price
    )

@router.post("/products/new")
async def product_create_save(
    request: Request,
    product_data: ProductCreateSchema = Depends(ProductCreateSchema.as_form),
    images: List[UploadFile] = File(default=[]),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "products")
    if denied: return denied
    embed_mode = _is_embed_request(request)

    category = await _get_category_for_assignment(session, product_data.category_id)
    if not category:
        return RedirectResponse(
            _build_admin_url("/admin/products/new", embed=embed_mode, error="invalid_category"),
            status_code=303,
        )
    if _has_invalid_product_discount(product_data):
        return RedirectResponse(
            _build_admin_url("/admin/products/new", embed=embed_mode, error="invalid_discount"),
            status_code=303,
        )

    saved_media = await save_uploaded_product_media_batch(images, limit=10)

    # Первое ФОТО — главное изображение товара (видео не может быть главным)
    main_image = ""
    for media in saved_media:
        if media.media_type == "image":
            main_image = media.public_path
            break

    new_product = Product(
        name_ru=product_data.name_ru,
        name_uz=product_data.name_uz,
        category_id=category.id,
        price=product_data.price,
        discount_price=product_data.discount_price,
        stock=product_data.stock,
        sell_type=product_data.sell_type or "piece",
        pack_quantity=product_data.pack_quantity if product_data.sell_type == "pack" else None,
        description_ru=product_data.description_ru,
        description_uz=product_data.description_uz,
        ikpu=product_data.ikpu,
        package_code=product_data.package_code,
        image_path=main_image,
        is_active=True
    )
    
    product_repo = ProductRepository(session)
    product_repo.add(new_product)
    
    try:
        await session.flush()  # получаем new_product.id
        # Сохраняем дополнительные медиа (пропускаем главное фото)
        sort_idx = 1
        for media in saved_media:
            if media.public_path == main_image:
                continue
            session.add(ProductImage(
                product_id=new_product.id,
                image_path=media.public_path,
                media_type=media.media_type,
                poster_path=media.poster_path,
                sort_order=sort_idx,
            ))
            sort_idx += 1
        await session.commit()
    except Exception:
        await session.rollback()
        for media in saved_media:
            await delete_product_media_files(media.public_path, media.poster_path)
        return RedirectResponse(
            _build_admin_url("/admin/products/new", embed=embed_mode, error="db_error"),
            status_code=303,
        )

    if embed_mode:
        return _modal_product_saved_response()
    return RedirectResponse("/admin/products", status_code=303)

@router.get("/products/{product_id}/edit", response_class=HTMLResponse)
async def product_edit_form(
    request: Request,
    product_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "products")
    if denied: return denied
    embed_mode = _is_embed_request(request)

    stmt = select(Product).options(selectinload(Product.images)).where(Product.id == product_id)
    product = (await session.execute(stmt)).scalar_one_or_none()
    
    if not product:
        return RedirectResponse("/admin/products")

    if not embed_mode:
        return RedirectResponse(f"/admin/products?modal=edit&product_id={product_id}", status_code=303)

    categories = await _get_assignable_categories(session, current_category_id=product.category_id)

    extra_images = sorted(product.images, key=lambda x: x.sort_order)

    return templates.TemplateResponse("admin/product_edit.html", {
        "request": request,
        "user": user,
        "categories": categories,
        "product": product,
        "extra_images": extra_images,
        "csrf_token": generate_csrf_token(request),
        "embed": embed_mode,
        "layout_template": "admin/embed_base.html" if embed_mode else "admin/base.html",
    })

@router.post("/products/{product_id}/edit")
async def product_edit_save(
    request: Request,
    product_id: int,
    product_data: ProductCreateSchema = Depends(ProductCreateSchema.as_form),
    images: List[UploadFile] = File(default=[]),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "products")
    if denied: return denied
    embed_mode = _is_embed_request(request)
    
    product_repo = ProductRepository(session)
    product = await product_repo.get_with_lock(product_id)
    
    if not product:
        return RedirectResponse("/admin/products")

    category = await _get_category_for_assignment(
        session,
        product_data.category_id,
        current_category_id=product.category_id,
    )
    # Запрещаем только перенос активного товара в архивную категорию.
    # Простое переименование (без смены категории) разрешаем всегда,
    # даже если товар оказался в архивной категории.
    category_changed = product_data.category_id != product.category_id
    if not category or (category_changed and not category.is_active and product.is_active):
        return RedirectResponse(
            _build_admin_url(f"/admin/products/{product_id}/edit", embed=embed_mode, error="invalid_category"),
            status_code=303,
        )
    if _has_invalid_product_discount(product_data):
        return RedirectResponse(
            _build_admin_url(f"/admin/products/{product_id}/edit", embed=embed_mode, error="invalid_discount"),
            status_code=303,
        )
        
    # Update fields
    was_out_of_stock = product.stock <= 0
    product.name_ru = product_data.name_ru
    product.name_uz = product_data.name_uz
    product.category_id = category.id
    product.price = product_data.price
    product.discount_price = product_data.discount_price
    product.stock = product_data.stock
    product.sell_type = product_data.sell_type or "piece"
    product.pack_quantity = product_data.pack_quantity if product_data.sell_type == "pack" else None
    product.description_ru = product_data.description_ru
    product.description_uz = product_data.description_uz
    product.ikpu = product_data.ikpu
    product.package_code = product_data.package_code

    saved_media = await save_uploaded_product_media_batch(images, limit=10)

    # If product has no main image, use first uploaded IMAGE as main
    if not product.image_path:
        for i, media in enumerate(saved_media):
            if media.media_type == "image":
                product.image_path = media.public_path
                saved_media.pop(i)
                break

    # Remaining uploaded media → ProductImage extras
    existing_max = 0
    if saved_media:
        max_sort = await session.execute(
            select(func.coalesce(func.max(ProductImage.sort_order), 0))
            .where(ProductImage.product_id == product_id)
        )
        existing_max = max_sort.scalar() or 0

    for idx, media in enumerate(saved_media, start=existing_max + 1):
        session.add(ProductImage(
            product_id=product_id,
            image_path=media.public_path,
            media_type=media.media_type,
            poster_path=media.poster_path,
            sort_order=idx,
        ))

    try:
        await session.commit()
    except Exception:
        await session.rollback()
        for media in saved_media:
            await delete_product_media_files(media.public_path, media.poster_path)
        return RedirectResponse(
            _build_admin_url(f"/admin/products/{product_id}/edit", embed=embed_mode, error="db_error"),
            status_code=303,
        )

    if was_out_of_stock and product_data.stock > 0:
        asyncio.create_task(_send_stock_notifications(product_id, product_data.name_ru))

    if embed_mode:
        return _modal_product_saved_response()
    return RedirectResponse("/admin/products", status_code=303)


@router.post("/products/popular/reorder", response_class=JSONResponse, dependencies=[Depends(validate_csrf_header)])
async def products_popular_reorder(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return JSONResponse({"error": "unauthorized"}, status_code=401)
    if not has_permission(user, "products"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    try:
        payload = await request.json()
        product_ids = [int(item) for item in payload.get("ids", [])]
    except (TypeError, ValueError, AttributeError):
        return JSONResponse({"error": "invalid_payload"}, status_code=400)

    if not product_ids:
        return JSONResponse({"error": "empty_order"}, status_code=400)
    if len(product_ids) != len(set(product_ids)):
        return JSONResponse({"error": "duplicate_products"}, status_code=400)

    products = (
        await session.execute(
            select(Product)
            .where(Product.id.in_(product_ids), Product.is_popular == True)
            .with_for_update()
        )
    ).scalars().all()
    product_by_id = {product.id: product for product in products}
    if len(product_by_id) != len(product_ids):
        return JSONResponse({"error": "product_not_found"}, status_code=404)

    for index, product_id in enumerate(product_ids, start=1):
        product_by_id[product_id].popular_sort_order = index * 10

    await session.commit()
    return JSONResponse({"ok": True})


@router.post("/products/{product_id}/popular")
async def product_update_popular(
    request: Request,
    product_id: int,
    action: str = Form(...),
    popular_sort_order: Optional[str] = Form("0"),
    return_to: str = Form("/admin/products?modal=popular"),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "products"):
        return RedirectResponse("/admin")

    redirect_url = _safe_products_return_url(return_to)
    product_repo = ProductRepository(session)
    product = await product_repo.get_with_lock(product_id)
    if not product:
        return RedirectResponse(redirect_url, status_code=303)

    normalized_action = (action or "").strip().lower()
    try:
        sort_order = max(int(popular_sort_order or 0), 0)
    except (ValueError, TypeError):
        sort_order = 0

    if normalized_action == "remove":
        product.is_popular = False
        product.popular_sort_order = 0
        await session.commit()
        return RedirectResponse(redirect_url, status_code=303)

    if normalized_action == "update_order":
        if product.is_popular:
            product.popular_sort_order = sort_order
            await session.commit()
        return RedirectResponse(redirect_url, status_code=303)

    if normalized_action == "add":
        category = (
            await session.execute(
                select(Category).where(Category.id == product.category_id)
            )
        ).scalar_one_or_none()
        if not product.is_active or not category or not category.is_active:
            return RedirectResponse(redirect_url, status_code=303)

        if sort_order <= 0:
            max_sort_order = (
                await session.execute(
                    select(func.coalesce(func.max(Product.popular_sort_order), 0))
                    .where(Product.is_popular == True)
                )
            ).scalar() or 0
            sort_order = int(max_sort_order) + 10

        product.is_popular = True
        product.popular_sort_order = sort_order
        await session.commit()

    return RedirectResponse(redirect_url, status_code=303)


@router.post("/products/{product_id}/toggle")
async def product_toggle_status(
    product_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "products"): return RedirectResponse("/admin")

    product_repo = ProductRepository(session)
    product = await product_repo.get_with_lock(product_id)

    if product:
        if not product.is_active:
            category = (await session.execute(
                select(Category).where(Category.id == product.category_id)
            )).scalar_one_or_none()
            if not category or not category.is_active:
                return RedirectResponse(
                    f"/admin/products?error={quote('Нельзя активировать товар, пока его категория находится в архиве.')}",
                    status_code=303,
                )
        was_active = product.is_active
        product.is_active = not product.is_active
        if was_active and not product.is_active:
            await session.execute(
                delete(CartItem).where(CartItem.product_id == product.id)
            )
        await session.commit()

    return RedirectResponse("/admin/products", status_code=303)

@router.post("/products/{product_id}/stock")
async def product_update_stock(
    product_id: int,
    stock: int = Form(...),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "products"): return RedirectResponse("/admin")

    if stock < 0:
        return RedirectResponse("/admin/products?error=invalid_stock", status_code=303)

    product_repo = ProductRepository(session)
    product = await product_repo.get_with_lock(product_id)

    if product:
        was_out_of_stock = product.stock <= 0
        product.stock = stock

        if stock == 0:
            await session.execute(
                delete(CartItem).where(CartItem.product_id == product_id)
            )

        await session.commit()

        if was_out_of_stock and stock > 0:
            asyncio.create_task(_send_stock_notifications(product_id, product.name_ru))

    return RedirectResponse("/admin/products", status_code=303)

@router.post("/products/delete/{product_id}")
async def product_delete(
    product_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    request: Request = None,
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "products"): return RedirectResponse("/admin")
    
    product_repo = ProductRepository(session)
    product = await product_repo.get_by_id(product_id)
    
    if product:
        # Soft Delete: помечаем как архивный вместо удаления
        # Это предотвращает ошибки связей с заказами и корзинами
        product.is_active = False
        await session.execute(
            delete(CartItem).where(CartItem.product_id == product.id)
        )
        await session.execute(
            delete(Favorite).where(Favorite.product_id == product.id)
        )
        await session.execute(
            delete(StockNotification).where(StockNotification.product_id == product.id)
        )
        # НЕ удаляем файл изображения — товар остается в истории заказов
        await session.commit()
        
    return RedirectResponse("/admin/products", status_code=303)

@router.get("/orders", response_class=HTMLResponse)
async def orders_list(
    request: Request,
    q: str = "",
    status: str = "all",
    payment: str = "all",
    order_type: str = "all",
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "orders")
    if denied: return denied

    page = 1
    try:
        page = int(request.query_params.get("page", 1))
    except (TypeError, ValueError):
        logger.opt(exception=True).debug("Invalid page param in orders list")
    if page < 1:
        page = 1
    
    limit = 20
    offset = (page - 1) * limit
    
    stmt = select(Order).options(selectinload(Order.user))
    if q:
        safe_query = q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        search_filters = [
            User.username.ilike(f"%{safe_query}%", escape="\\"),
            User.phone.ilike(f"%{safe_query}%", escape="\\"),
            Order.contact_phone.ilike(f"%{safe_query}%", escape="\\"),
            Order.customer_name.ilike(f"%{safe_query}%", escape="\\"),
        ]
        if q.isdigit() and len(q) <= 9:
            int_q = int(q)
            if int_q <= 2147483647:
                search_filters.append(Order.id == int_q)
        stmt = stmt.outerjoin(User).where(or_(*search_filters))

    if status != "all":
        stmt = stmt.where(Order.status == status)
    if payment != "all":
        stmt = stmt.where(Order.payment_method == payment)
    if order_type != "all":
        stmt = stmt.where(Order.order_type == order_type)

    total_count_stmt = select(func.count()).select_from(stmt.subquery())
    total_count = (await session.execute(total_count_stmt)).scalar() or 0

    orders = (await session.execute(
        stmt.order_by(Order.created_at.desc()).limit(limit).offset(offset)
    )).scalars().all()
    
    total_pages = (total_count + limit - 1) // limit

    status_counts_stmt = (
        select(Order.status, func.count(Order.id))
        .group_by(Order.status)
    )
    status_counts_raw = (await session.execute(status_counts_stmt)).all()
    status_counts = {row[0]: row[1] for row in status_counts_raw}

    revenue_stmt = select(func.sum(Order.total_amount)).where(Order.status.in_(['done', 'paid']))
    revenue_total = (await session.execute(revenue_stmt)).scalar() or 0

    return templates.TemplateResponse("admin/orders_list.html", {
        "request": request, 
        "user": user, 
        "orders": orders,
        "page": page,
        "total_pages": total_pages,
        "filters": {"q": q, "status": status, "payment": payment, "order_type": order_type},
        "status_counts": status_counts,
        "revenue_total": revenue_total,
        "csrf_token": generate_csrf_token(request)
    })


@router.get("/orders/unread-count")
async def admin_orders_unread_count(
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return JSONResponse({"error": "unauthorized"}, status_code=401)
    if not has_permission(user, "orders"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    return JSONResponse({"count": await _count_unread_new_orders(session)})

# ===================== СОЗДАНИЕ ЗАКАЗА ВРУЧНУЮ =====================
@router.get("/api/products/search", response_class=JSONResponse)
async def api_products_search(
    request: Request,
    q: str = "",
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    """API endpoint for searching products (used in order creation form)."""
    if not user:
        return JSONResponse({"error": "unauthorized"}, status_code=401)
    if not has_permission(user, "orders"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    stmt = select(Product).where(Product.is_active == True)
    if q:
        safe_query = q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        stmt = stmt.where(
            or_(
                Product.name_ru.ilike(f"%{safe_query}%", escape="\\"),
                Product.name_uz.ilike(f"%{safe_query}%", escape="\\")
            )
        )
    stmt = stmt.order_by(Product.name_ru).limit(30)
    products = (await session.execute(stmt)).scalars().all()

    return JSONResponse([{
        "id": p.id,
        "name": p.name_ru,
        "price": p.effective_price,
        "regular_price": p.price,
        "discount_price": p.discount_price if p.has_discount else None,
        "has_discount": p.has_discount,
        "discount_percent": p.discount_percent,
        "stock": p.stock,
        "image": p.image_path or "",
        "sell_type": p.sell_type or "piece",
        "pack_quantity": p.pack_quantity,
        "min_order_quantity": get_min_order_quantity_for_product(p),
    } for p in products])


@router.get("/orders/new", response_class=HTMLResponse)
async def order_create_form(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "orders")
    if denied:
        return denied

    return templates.TemplateResponse("admin/order_create.html", {
        "request": request,
        "user": user,
        "csrf_token": generate_csrf_token(request),
        "pickup_enabled": await get_pickup_enabled(session),
        "pickup_address": await get_pickup_address(session),
        "delivery_options": await get_delivery_options(session, active_only=True),
    })


@router.post("/orders/new")
async def order_create_save(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "orders")
    if denied:
        return denied

    form = await request.form()
    customer_name = (form.get("customer_name") or "").strip()
    contact_phone = (form.get("contact_phone") or "").strip()
    payment_method = form.get("payment_method", "cash")
    delivery_method = form.get("delivery_method", "pickup")
    delivery_option_id = form.get("delivery_option_id")
    delivery_address = (form.get("delivery_address") or "").strip()
    comment = (form.get("comment") or "").strip()
    pickup_enabled = await get_pickup_enabled(session)
    pickup_address = await get_pickup_address(session)
    delivery_options = await get_delivery_options(session, active_only=True)
    selected_delivery_option = find_delivery_option(delivery_options, delivery_option_id)
    delivery_option_name = None
    delivery_price = 0
    if delivery_method == "delivery" and selected_delivery_option:
        delivery_option_id = selected_delivery_option["id"]
        delivery_option_name = selected_delivery_option.get("name_ru") or selected_delivery_option["id"]
        delivery_price = int(selected_delivery_option.get("price") or 0)
    form_context = {
        "pickup_enabled": pickup_enabled,
        "pickup_address": pickup_address,
        "delivery_options": delivery_options,
    }

    # Validate
    errors = []
    if not customer_name:
        errors.append("Укажите имя клиента")
    if not contact_phone or len(contact_phone) < 9:
        errors.append("Укажите корректный номер телефона")
    if payment_method not in ("cash", "card", "click", "card_transfer"):
        errors.append("Неверный способ оплаты")
    if delivery_method not in ("pickup", "delivery"):
        errors.append("Неверный способ доставки")
    if delivery_method == "pickup" and not pickup_enabled:
        errors.append("Самовывоз отключён в настройках")
    if delivery_method == "delivery" and not delivery_address:
        errors.append("Укажите адрес доставки")
    if delivery_method == "delivery" and not selected_delivery_option:
        errors.append("Выберите вариант доставки")

    # Parse product items from form
    product_ids = form.getlist("product_id[]")
    quantities = form.getlist("quantity[]")

    if not product_ids:
        errors.append("Добавьте хотя бы один товар")

    if errors:
        return templates.TemplateResponse("admin/order_create.html", {
            "request": request,
            "user": user,
            "csrf_token": generate_csrf_token(request),
            "errors": errors,
            "form_data": {
                "customer_name": customer_name,
                "contact_phone": contact_phone,
                "payment_method": payment_method,
                "delivery_method": delivery_method,
                "delivery_option_id": delivery_option_id,
                "delivery_address": delivery_address,
                "comment": comment,
            },
            **form_context,
        })

    # Process products and create order
    try:
        total_amount = 0
        items_to_create = []

        for pid_str, qty_str in zip(product_ids, quantities):
            try:
                pid = int(pid_str)
                qty = int(qty_str)
            except (ValueError, TypeError):
                continue
            if qty < 1:
                continue

            # Atomic stock decrement
            product = await session.get(Product, pid)
            if not product or not product.is_active:
                errors.append(f"Товар ID {pid} не найден или неактивен")
                continue
            min_order_quantity = get_min_order_quantity_for_product(product)

            if qty < min_order_quantity:
                errors.append(f"Минимальное количество для товара \"{product.name_ru}\" — {min_order_quantity} шт")
                continue

            if product.stock < qty:
                errors.append(f"Недостаточно товара \"{product.name_ru}\" (на складе: {product.stock})")
                continue

            stock_before = product.stock  # capture BEFORE atomic decrement

            from sqlalchemy import update as sa_update
            result = await session.execute(
                sa_update(Product)
                .where(Product.id == pid, Product.stock >= qty)
                .values(stock=Product.stock - qty)
                .execution_options(synchronize_session="fetch")
            )
            if result.rowcount == 0:
                errors.append(f"Не удалось списать товар \"{product.name_ru}\"")
                continue

            unit_price = product.effective_price
            item_total = unit_price * qty
            total_amount += item_total
            p_name = product.name_ru
            if product.sell_type == "pack" and product.pack_quantity:
                p_name += f" (уп. {product.pack_quantity} шт)"
            items_to_create.append({
                "product_id": product.id,
                "product_name": p_name,
                "price_at_purchase": unit_price,
                "quantity": qty,
                "stock_before_order": stock_before,
            })

        if errors or not items_to_create:
            await session.rollback()
            if not items_to_create and not errors:
                errors.append("Не удалось добавить ни одного товара")
            return templates.TemplateResponse("admin/order_create.html", {
                "request": request,
                "user": user,
                "csrf_token": generate_csrf_token(request),
                "errors": errors,
                "form_data": {
                    "customer_name": customer_name,
                    "contact_phone": contact_phone,
                    "payment_method": payment_method,
                    "delivery_method": delivery_method,
                    "delivery_option_id": delivery_option_id,
                    "delivery_address": delivery_address,
                    "comment": comment,
                },
                **form_context,
            })

        total_amount += delivery_price

        # Create order without user_id
        new_order = Order(
            user_id=None,
            customer_name=customer_name,
            status="new",
            admin_unread=False,
            order_type="product",
            payment_method=payment_method,
            delivery_method=delivery_method,
            delivery_option_id=delivery_option_id if delivery_method == "delivery" else None,
            delivery_option_name=delivery_option_name if delivery_method == "delivery" else None,
            delivery_price=delivery_price if delivery_method == "delivery" else 0,
            delivery_address=delivery_address if delivery_method == "delivery" else pickup_address,
            total_amount=total_amount,
            discount_amount=0,
            comment=comment,
            contact_phone=contact_phone,
        )
        session.add(new_order)
        await session.flush()

        for item_data in items_to_create:
            session.add(OrderItem(
                order_id=new_order.id,
                **item_data,
            ))

        await log_action(
            session, "Создание заказа вручную", request=request,
            user_id=user.id, entity_type="заказ", entity_id=new_order.id,
            details={"клиент": customer_name, "телефон": contact_phone, "сумма": total_amount}
        )
        await session.commit()

        return RedirectResponse(f"/admin/orders/{new_order.id}", status_code=303)

    except Exception as exc:
        await session.rollback()
        logger.exception("Failed to create manual order")
        return templates.TemplateResponse("admin/order_create.html", {
            "request": request,
            "user": user,
            "csrf_token": generate_csrf_token(request),
            "errors": [f"Ошибка создания заказа: {exc}"],
            "form_data": {
                "customer_name": customer_name,
                "contact_phone": contact_phone,
                "payment_method": payment_method,
                "delivery_method": delivery_method,
                "delivery_option_id": delivery_option_id,
                "delivery_address": delivery_address,
                "comment": comment,
            },
            **form_context,
        })


# ===================== ЭКСПОРТ ЗАКАЗОВ XLSX =====================
@router.get("/orders/export")
async def orders_export_xlsx(
    request: Request,
    date_from: str = "",
    date_to: str = "",
    status: str = "all",
    payment: str = "all",
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "orders"): return RedirectResponse("/admin")

    stmt = (
        select(Order)
        .options(selectinload(Order.user), selectinload(Order.items))
        .order_by(Order.created_at.desc())
    )

    if date_from:
        try:
            dt_from = datetime.strptime(date_from, "%Y-%m-%d")
            stmt = stmt.where(Order.created_at >= dt_from)
        except ValueError:
            pass
    if date_to:
        try:
            dt_to = datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1)
            stmt = stmt.where(Order.created_at < dt_to)
        except ValueError:
            pass
    if status != "all":
        stmt = stmt.where(Order.status == status)
    if payment != "all":
        stmt = stmt.where(Order.payment_method == payment)

    stmt = stmt.limit(50000)
    orders = (await session.execute(stmt)).scalars().all()

    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заказы"

    headers = ["ID", "Дата", "Клиент", "Телефон", "Статус", "Оплата", "Доставка", "Стоимость доставки", "Адрес", "Геолокация", "Сумма", "Скидка", "Товары"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="3E2310", end_color="3E2310", fill_type="solid")
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    xl_status = {"new": "Новый", "paid": "Оплачен", "delivery": "В пути", "done": "Выполнен", "cancelled": "Отменён"}
    xl_pay = {"cash": "Наличные", "card": "Payme", "click": "Click", "card_transfer": "Перевод на карту"}
    xl_delivery = {"delivery": "Доставка", "pickup": "Самовывоз"}

    for row_idx, o in enumerate(orders, 2):
        client_name = o.user.username if o.user else (o.customer_name or "—")
        items_str = "; ".join(f"{i.product_name} x{i.quantity}" for i in o.items)
        ws.cell(row=row_idx, column=1, value=o.id)
        ws.cell(row=row_idx, column=2, value=format_datetime_uz(o.created_at))
        ws.cell(row=row_idx, column=3, value=client_name)
        ws.cell(row=row_idx, column=4, value=o.contact_phone)
        ws.cell(row=row_idx, column=5, value=xl_status.get(o.status, o.status))
        ws.cell(row=row_idx, column=6, value=xl_pay.get(o.payment_method, o.payment_method))
        delivery_label = "Самовывоз" if o.delivery_method == "pickup" else (o.delivery_option_name or xl_delivery.get(o.delivery_method, o.delivery_method or ""))
        ws.cell(row=row_idx, column=7, value=delivery_label)
        ws.cell(row=row_idx, column=8, value=o.delivery_price or 0)
        ws.cell(row=row_idx, column=9, value=o.delivery_address or "")
        location = (
            f"{o.delivery_latitude:.6f}, {o.delivery_longitude:.6f}"
            if o.delivery_latitude is not None and o.delivery_longitude is not None
            else ""
        )
        ws.cell(row=row_idx, column=10, value=location)
        ws.cell(row=row_idx, column=11, value=o.total_amount)
        ws.cell(row=row_idx, column=12, value=o.discount_amount)
        ws.cell(row=row_idx, column=13, value=items_str)

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    await log_action(session, "Экспорт заказов", request=request, user_id=user.id)
    await session.commit()

    from starlette.responses import StreamingResponse
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=orders_export.xlsx"},
    )


@router.get("/orders/{order_id}/receipt")
async def order_payment_receipt(
    order_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        raise HTTPException(status_code=401, detail="Unauthorized")
    if not has_permission(user, "orders"):
        raise HTTPException(status_code=403, detail="Forbidden")

    order = await session.get(Order, order_id)
    if not order or not order.payment_receipt_path:
        raise HTTPException(status_code=404, detail="Receipt not found")

    receipt_path = public_media_to_path(order.payment_receipt_path)
    if not receipt_path or not receipt_path.is_file():
        raise HTTPException(status_code=404, detail="Receipt not found")

    return FileResponse(receipt_path)


@router.get("/orders/{order_id}", response_class=HTMLResponse)
async def order_detail(
    request: Request,
    order_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "orders")
    if denied: return denied

    await session.execute(
        update(Order)
        .where(Order.id == order_id, Order.admin_unread == True)
        .values(admin_unread=False)
    )
    await session.commit()

    order_repo = OrderRepository(session)
    order = await order_repo.get_full_info(order_id)
    
    if not order:
        return RedirectResponse("/admin/orders")

    error_message = request.query_params.get("error")
    if error_message:
        error_message = unquote(error_message)
    success_message = request.query_params.get("success")
    if success_message:
        success_message = unquote(success_message)

    # Generate payment URL for unpaid online orders (so admin can share with client)
    pay_url = None
    if order.status == "new" and order.payment_method in ("card", "click"):
        from app.utils.payment import generate_payme_link, generate_click_link
        if order.payment_method == "card":
            pay_url = generate_payme_link(order.id, order.total_amount)
        else:
            pay_url = generate_click_link(order.id, order.total_amount)

    return templates.TemplateResponse("admin/order_detail.html", {
        "request": request, 
        "user": user, 
        "order": order,
        "error": error_message,
        "success": success_message,
        "pay_url": pay_url,
        "csrf_token": generate_csrf_token(request)
    })

@router.post("/orders/{order_id}/delete")
async def order_delete(
    order_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    request: Request = None,
    csrf: bool = Depends(validate_csrf)
):
    if not user:
        return RedirectResponse("/admin/login")
    if user.role != "superadmin":
        return RedirectResponse("/admin", status_code=303)

    order_repo = OrderRepository(session)
    order = await order_repo.get_with_lock(order_id)

    if not order:
        return RedirectResponse("/admin/orders", status_code=303)

    if order.user_id:
        error_message = "Удалять можно только вручную созданные заказы."
        return RedirectResponse(
            f"/admin/orders/{order_id}?error={quote(error_message)}",
            status_code=303,
        )

    if order.status not in ("new", "cancelled"):
        error_message = "Удалять можно только новые или отменённые заказы."
        return RedirectResponse(
            f"/admin/orders/{order_id}?error={quote(error_message)}",
            status_code=303,
        )

    await OrderService.cancel_order(session, order_id, commit=False)

    await session.execute(delete(PaymeTransaction).where(PaymeTransaction.order_id == order_id))
    await session.execute(delete(OrderItem).where(OrderItem.order_id == order_id))
    if order.payment_receipt_path:
        await delete_file(order.payment_receipt_path)
    await session.execute(delete(Order).where(Order.id == order_id))
    await session.execute(delete(ClickTransaction).where(ClickTransaction.merchant_trans_id == str(order_id)))

    await log_action(
        session,
        "Удаление заказа",
        request=request,
        user_id=user.id,
        entity_type="заказ",
        entity_id=order_id,
        details={"manual": True},
    )
    await session.commit()

    return RedirectResponse("/admin/orders", status_code=303)

@router.post("/orders/{order_id}/status")
async def order_change_status(
    order_id: int,
    status: str = Form(...),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    request: Request = None,
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "orders"): return RedirectResponse("/admin")

    order_repo = OrderRepository(session)
    online_payment_methods = ("card", "click")

    order = await order_repo.get_with_lock(order_id)
    if not order:
        return RedirectResponse("/admin/orders", status_code=303)

    if order.status == "cancelled":
        error_message = "Отмененный заказ нельзя изменять."
        return RedirectResponse(
            f"/admin/orders/{order_id}?error={quote(error_message)}",
            status_code=303,
        )

    if order.status == "done" and status in ("new", "delivery", "paid"):
        error_message = (
            "Завершенный заказ нельзя переводить обратно."
        )
        return RedirectResponse(
            f"/admin/orders/{order_id}?error={quote(error_message)}",
            status_code=303,
        )

    if order.status in ("paid", "delivery") and status == "new":
        error_message = (
            "Нельзя переводить оплаченный или доставляемый заказ обратно в новый."
        )
        return RedirectResponse(
            f"/admin/orders/{order_id}?error={quote(error_message)}",
            status_code=303,
        )

    if status == "cancelled":
        if order.status in ("paid", "done", "delivery"):
            error_message = (
                "Нельзя отменять оплаченный, доставляемый или завершенный заказ — "
                "возвраты через админку не поддерживаются."
            )
            return RedirectResponse(
                f"/admin/orders/{order_id}?error={quote(error_message)}",
                status_code=303,
            )
        order = await OrderService.cancel_order(session, order_id, commit=False)
    else:
        if status in ("delivery", "done") and order.status == "new":
            error_message = None
            if order.payment_method in online_payment_methods:
                error_message = "Онлайн-заказ нельзя отправлять в доставку или завершать до оплаты."
            elif order.payment_method == "card_transfer":
                error_message = "Сначала проверьте чек и подтвердите оплату."
            if error_message:
                return RedirectResponse(
                    f"/admin/orders/{order_id}?error={quote(error_message)}",
                    status_code=303,
                )
        order.status = status
    
    if order.status != "new":
        order.admin_unread = False

    # Audit log + single commit
    status_names_ru = {"new": "Новый", "paid": "Оплачен", "delivery": "В пути", "done": "Выполнен", "cancelled": "Отменён"}
    await log_action(session, "Смена статуса заказа", request=request, user_id=user.id,
                     entity_type="заказ", entity_id=order_id, details={"новый_статус": status_names_ru.get(status, status)})
    await session.commit()
        
    if order:
        OrderService.notify_user_order_status(order.user, status)
        OrderService.notify_status_subscribers(order_id, status)

    return RedirectResponse(f"/admin/orders/{order_id}", status_code=303)


@router.post("/orders/{order_id}/restore")
async def order_restore_cancelled_cash(
    order_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    request: Request = None,
    csrf: bool = Depends(validate_csrf),
):
    """Restore only cancelled cash product orders from the admin panel."""
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "orders"):
        return RedirectResponse("/admin", status_code=303)

    try:
        order = await OrderService.restore_cancelled_cash_order(session, order_id, commit=False)
        if not order:
            return RedirectResponse("/admin/orders", status_code=303)

        await log_action(
            session,
            "Восстановление отменённого заказа",
            request=request,
            user_id=user.id,
            entity_type="заказ",
            entity_id=order_id,
            details={"способ_оплаты": "Наличные", "новый_статус": "Новый"},
        )
        await session.commit()
    except OrderRestoreError as exc:
        await session.rollback()
        return RedirectResponse(
            f"/admin/orders/{order_id}?error={quote(str(exc))}",
            status_code=303,
        )
    except IntegrityError:
        await session.rollback()
        return RedirectResponse(
            f"/admin/orders/{order_id}?error={quote('Данные заказа изменились. Обновите страницу и попробуйте ещё раз.')}",
            status_code=303,
        )

    OrderService.notify_user_order_status(order.user, "restored")
    OrderService.notify_status_subscribers(order_id, "new")
    return RedirectResponse(
        f"/admin/orders/{order_id}?success={quote('Заказ восстановлен и снова ожидает обработки.')}",
        status_code=303,
    )

@router.get("/users", response_class=HTMLResponse)
async def users_list(
    request: Request,
    q: str = "",
    debt: str = "all",
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "users")
    if denied: return denied

    stmt = select(User).where(User.role == "user")
    if q:
        safe_query = q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        search_filters = [
            User.username.ilike(f"%{safe_query}%", escape="\\"),
            User.phone.ilike(f"%{safe_query}%", escape="\\")
        ]
        if q.isdigit() and len(q) <= 9:
            int_q = int(q)
            if int_q <= 2147483647:
                search_filters.append(User.id == int_q)
        stmt = stmt.where(or_(*search_filters))

    if debt == "with":
        stmt = stmt.where(User.debt > 0)
    elif debt == "without":
        stmt = stmt.where(or_(User.debt == 0, User.debt.is_(None)))

    users = (await session.execute(stmt.order_by(User.id.desc()))).scalars().all()

    return templates.TemplateResponse("admin/users_list.html", {
        "request": request,
        "user": user,
        "users": users,
        "filters": {"q": q, "debt": debt},
        "csrf_token": generate_csrf_token(request)
    })

@router.post("/users/{user_id}/set_debt")
async def user_set_debt(
    user_id: int,
    amount: Optional[int] = Form(None),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "users"): return RedirectResponse("/admin")

    if amount is None:
        amount = 0
    if amount < 0:
        return RedirectResponse("/admin/users?error=invalid_debt", status_code=303)
    
    repo = UserRepository(session)
    target_user = await repo.get_with_lock(user_id)
    
    if target_user:
        target_user.debt = amount
        await session.commit()
        
    return RedirectResponse("/admin/users", status_code=303)

@router.get("/managers", response_class=HTMLResponse)
async def managers_list(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user:
        return RedirectResponse("/admin/login")
    if user.role != "superadmin":
        return templates.TemplateResponse("admin/error.html", {"request": request, "message": "Доступ запрещен"})

    repo = UserRepository(session)
    managers = await repo.get_admins()
    for manager in managers:
        if manager.role == "manager":
            manager.permissions = normalize_permissions(manager.permissions)
    error = request.query_params.get("error")

    return templates.TemplateResponse("admin/managers_list.html", {
        "request": request, 
        "user": user, 
        "managers": managers,
        "sections": ADMIN_SECTIONS,
        "error": error,
        "csrf_token": generate_csrf_token(request)
    })

@router.post("/managers/new")
async def manager_create(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    telegram_id: int = Form(None),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user:
        return RedirectResponse("/admin/login")
    if user.role != "superadmin":
        return RedirectResponse("/admin")
    
    from app.utils.security import get_password_hash
    pwd_hash = get_password_hash(password)
    
    if not telegram_id:
        telegram_id = None

    # Collect permissions from form checkboxes
    form_data = await request.form()
    perms = {}
    for section_key in ADMIN_SECTIONS:
        perms[section_key] = f"perm_{section_key}" in form_data
    perms = normalize_permissions(perms)

    new_manager = User(
        username="Менеджер",
        login=username,
        password_hash=pwd_hash,
        permissions=perms,
        role="manager",
        telegram_id=telegram_id
    )
    
    try:
        session.add(new_manager)
        await session.commit()
    except Exception as e:
        logger.exception("Failed to create manager")
        await session.rollback()
        error_message = "Не удалось создать менеджера (возможно, логин уже занят)"
        return RedirectResponse(
            f"/admin/managers?error={quote(error_message)}",
            status_code=303
        )

    return RedirectResponse("/admin/managers", status_code=303)

@router.get("/managers/{manager_id}", response_class=HTMLResponse)
async def manager_edit_page(
    request: Request,
    manager_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user:
        return RedirectResponse("/admin/login")
    if user.role != "superadmin":
        return templates.TemplateResponse("admin/error.html", {"request": request, "message": "Доступ запрещён"})

    manager = await session.get(User, manager_id)
    if not manager or manager.role not in ("manager",):
        return RedirectResponse("/admin/managers")
    manager.permissions = normalize_permissions(manager.permissions)

    error = request.query_params.get("error")
    success = request.query_params.get("success")

    return templates.TemplateResponse("admin/manager_edit.html", {
        "request": request,
        "user": user,
        "manager": manager,
        "sections": ADMIN_SECTIONS,
        "error": error,
        "success": success,
        "csrf_token": generate_csrf_token(request),
    })

@router.post("/managers/{manager_id}/permissions")
async def manager_update_permissions(
    request: Request,
    manager_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if user.role != "superadmin":
        return RedirectResponse("/admin")

    manager = await session.get(User, manager_id)
    if not manager or manager.role != "manager":
        return RedirectResponse("/admin/managers")

    form_data = await request.form()
    perms = {}
    for section_key in ADMIN_SECTIONS:
        perms[section_key] = f"perm_{section_key}" in form_data
    perms = normalize_permissions(perms)

    # Explicit SQL UPDATE — bypasses ORM identity-map issues
    await session.execute(
        update(User).where(User.id == manager_id).values(permissions=perms)
    )
    await session.commit()
    logger.info("Permissions updated for manager %s: %s", manager_id, perms)

    return RedirectResponse(f"/admin/managers/{manager_id}?success={quote('Разрешения обновлены')}", status_code=303)

@router.post("/managers/{manager_id}/credentials")
async def manager_update_credentials(
    request: Request,
    manager_id: int,
    new_login: str = Form(...),
    new_password: str = Form(""),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if user.role != "superadmin":
        return RedirectResponse("/admin")

    manager = await session.get(User, manager_id)
    if not manager or manager.role != "manager":
        return RedirectResponse("/admin/managers")

    # Check login uniqueness
    if new_login != manager.login:
        existing = await UserRepository(session).get_by_login(new_login)
        if existing:
            return RedirectResponse(f"/admin/managers/{manager_id}?error={quote('Логин уже занят')}", status_code=303)
        manager.login = new_login

    if new_password.strip():
        from app.utils.security import get_password_hash
        manager.password_hash = get_password_hash(new_password.strip())

    await session.commit()
    return RedirectResponse(f"/admin/managers/{manager_id}?success={quote('Данные обновлены')}", status_code=303)

@router.post("/managers/delete/{manager_id}")
async def manager_delete(
    manager_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    request: Request = None,
    csrf: bool = Depends(validate_csrf)
):
    if not user:
        return RedirectResponse("/admin/login")
    if user.role != "superadmin":
        return RedirectResponse("/admin")

    repo = UserRepository(session)
    manager = await repo.get_by_id(manager_id)
    
    if manager and manager.role == "manager":
        await session.delete(manager)
        await session.commit()
        
    return RedirectResponse("/admin/managers", status_code=303)

async def perform_mailing(chat_ids: List[int], text: str, photo_bytes: Optional[bytes]):
    file_id = None

    # Pre-upload photo to Telegram to get a reliable file_id before the broadcast loop.
    # This avoids re-uploading bytes on every failed attempt when early users blocked the bot.
    if photo_bytes:
        from app.config import settings as _settings
        pre_upload_targets = list(_settings.ADMIN_IDS or [])
        for target_id in pre_upload_targets:
            try:
                pre_msg = await bot.send_photo(
                    target_id,
                    photo=BufferedInputFile(photo_bytes, filename="mailing.jpg"),
                    caption="[Тест рассылки — можно удалить]",
                )
                file_id = pre_msg.photo[-1].file_id
                try:
                    await bot.delete_message(target_id, pre_msg.message_id)
                except Exception:
                    pass
                break
            except Exception:
                continue

    for chat_id in chat_ids:
        try:
            if photo_bytes:
                if file_id is None:
                    # Fallback: upload via first successful send
                    inp = BufferedInputFile(photo_bytes, filename="mailing.jpg")
                    msg = await bot.send_photo(chat_id, photo=inp, caption=text, parse_mode="HTML")
                    file_id = msg.photo[-1].file_id
                else:
                    await bot.send_photo(chat_id, photo=file_id, caption=text, parse_mode="HTML")
            else:
                await bot.send_message(chat_id, text, parse_mode="HTML")
            await asyncio.sleep(0.05)
        except Exception as e:
            # Пользователь мог заблокировать бота
            logger.opt(exception=True).info("Failed to send mailing message")

@router.get("/mailing", response_class=HTMLResponse)
async def mailing_page(request: Request, user: User = Depends(get_current_admin)):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "mailing")
    if denied: return denied
    return templates.TemplateResponse("admin/mailing.html", {"request": request, "user": user, "csrf_token": generate_csrf_token(request)})

@router.post("/mailing/send")
async def mailing_send(
    request: Request,
    background_tasks: BackgroundTasks,
    text: str = Form(...),
    image: UploadFile = File(None),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "mailing"): return RedirectResponse("/admin")

    stmt = select(User.telegram_id).where(
        User.telegram_id.isnot(None),
        User.role == "user"
    )
    ids = (await session.execute(stmt)).scalars().all()

    photo_bytes = None
    if image and image.filename:
        try:
            photo_bytes = await read_upload_limited(image, IMAGE_MAX_BYTES)
            validate_image_upload(image.filename, photo_bytes)
        except Exception:
            photo_bytes = None # Skip invalid image

    background_tasks.add_task(perform_mailing, ids, text, photo_bytes)

    return templates.TemplateResponse("admin/mailing.html", {
        "request": request, 
        "user": user, 
        "message": f"Рассылка запущена для {len(ids)} пользователей. Она будет выполнена в фоновом режиме.",
        "csrf_token": generate_csrf_token(request)
    })

# ===================== АУДИТ ЛОГ =====================
@router.get("/audit", response_class=HTMLResponse)
async def audit_log_page(
    request: Request,
    page: int = 1,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "audit")
    if denied: return denied

    limit = 50
    offset = (max(page, 1) - 1) * limit

    total = (await session.execute(select(func.count(AuditLog.id)))).scalar() or 0
    stmt = (
        select(AuditLog)
        .options(selectinload(AuditLog.user))
        .order_by(AuditLog.created_at.desc())
        .limit(limit).offset(offset)
    )
    logs = (await session.execute(stmt)).scalars().all()
    total_pages = (total + limit - 1) // limit

    return templates.TemplateResponse("admin/audit_log.html", {
        "request": request, "user": user, "logs": logs,
        "page": page, "total_pages": total_pages,
        "csrf_token": generate_csrf_token(request)
    })

# ===================== ПРОМОКОДЫ =====================
@router.get("/promo", response_class=HTMLResponse)
async def promo_list(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "promo")
    if denied: return denied

    promos = (await session.execute(
        select(PromoCode).order_by(PromoCode.created_at.desc())
    )).scalars().all()

    return templates.TemplateResponse("admin/promo_list.html", {
        "request": request, "user": user, "promos": promos,
        "csrf_token": generate_csrf_token(request)
    })

@router.post("/promo/new")
async def promo_create(
    request: Request,
    code: str = Form(...),
    discount_type: str = Form("percent"),
    discount_value: int = Form(...),
    min_order_amount: int = Form(0),
    max_discount: Optional[int] = Form(None),
    usage_limit: Optional[int] = Form(None),
    expires_at: Optional[str] = Form(None),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "promo"): return RedirectResponse("/admin")

    if discount_type not in ("percent", "fixed"):
        return RedirectResponse("/admin/promo?error=invalid_type", status_code=303)
    if discount_value <= 0:
        return RedirectResponse("/admin/promo?error=invalid_value", status_code=303)
    if discount_type == "percent" and discount_value > 100:
        return RedirectResponse("/admin/promo?error=invalid_percent", status_code=303)

    exp_dt = None
    if expires_at:
        try:
            exp_dt = datetime.strptime(expires_at, "%Y-%m-%dT%H:%M")
            exp_dt = exp_dt - timedelta(hours=5)
        except ValueError:
            try:
                exp_dt = datetime.strptime(expires_at, "%Y-%m-%d")
                exp_dt = exp_dt - timedelta(hours=5)
            except ValueError:
                pass

    promo = PromoCode(
        code=code.strip().upper(),
        discount_type=discount_type,
        discount_value=discount_value,
        min_order_amount=min_order_amount,
        max_discount=max_discount,
        usage_limit=usage_limit,
        expires_at=exp_dt,
    )
    session.add(promo)
    await log_action(session, "Создание промокода", request=request, user_id=user.id, entity_type="промокод", details={"код": promo.code})
    try:
        await session.commit()
    except Exception:
        await session.rollback()
        return RedirectResponse("/admin/promo?error=duplicate", status_code=303)

    return RedirectResponse("/admin/promo", status_code=303)

@router.post("/promo/{promo_id}/toggle")
async def promo_toggle(
    promo_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "promo"): return RedirectResponse("/admin")
    promo = (await session.execute(select(PromoCode).where(PromoCode.id == promo_id))).scalar_one_or_none()
    if promo:
        promo.is_active = not promo.is_active
        await session.commit()
    return RedirectResponse("/admin/promo", status_code=303)

@router.post("/promo/{promo_id}/delete")
async def promo_delete(
    promo_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "promo"): return RedirectResponse("/admin")
    promo = (await session.execute(select(PromoCode).where(PromoCode.id == promo_id))).scalar_one_or_none()
    if promo:
        await session.delete(promo)
        await log_action(session, "Удаление промокода", request=None, user_id=user.id, entity_type="промокод", entity_id=promo_id)
        try:
            await session.commit()
        except IntegrityError:
            await session.rollback()
            promo = (await session.execute(select(PromoCode).where(PromoCode.id == promo_id))).scalar_one_or_none()
            if promo:
                promo.is_active = False
                await session.commit()
            return RedirectResponse("/admin/promo?error=has_orders", status_code=303)
    return RedirectResponse("/admin/promo", status_code=303)

# ===================== ПРОМО-БАННЕРЫ =====================
@router.get("/banners", response_class=HTMLResponse)
async def banners_list(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db)
):
    if not user: return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "banners")
    if denied: return denied

    banners = (await session.execute(
        select(PromoBanner)
        .options(selectinload(PromoBanner.products))
        .order_by(PromoBanner.sort_order.asc(), PromoBanner.id.desc())
    )).scalars().all()

    return templates.TemplateResponse("admin/banners_list.html", {
        "request": request, "user": user, "banners": banners,
        "csrf_token": generate_csrf_token(request)
    })

@router.post("/banners/new")
async def banner_create(
    request: Request,
    title: str = Form(""),
    subtitle: str = Form(""),
    description: str = Form(""),
    badge_text: str = Form(""),
    link: str = Form(""),
    sort_order: int = Form(0),
    image: UploadFile = File(...),
    image_crop_x: Optional[float] = Form(None),
    image_crop_y: Optional[float] = Form(None),
    image_crop_width: Optional[float] = Form(None),
    image_crop_height: Optional[float] = Form(None),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "banners"): return RedirectResponse("/admin")

    if not image or not image.filename:
        return RedirectResponse("/admin/banners?error=no_image", status_code=303)

    media_crop = _category_crop_payload(
        image_crop_x,
        image_crop_y,
        image_crop_width,
        image_crop_height,
    )
    image_path = await _save_banner_media(image, crop=media_crop)
    if image and image.filename and not image_path:
        return RedirectResponse("/admin/banners?error=invalid_image", status_code=303)

    banner = PromoBanner(
        title=title or None, subtitle=subtitle or None, description=description or None,
        badge_text=badge_text or None, image_path=image_path or None,
        link=link or None, sort_order=sort_order,
    )
    session.add(banner)
    await log_action(session, "Создание баннера", request=request, user_id=user.id, entity_type="баннер", details={"заголовок": title})
    await session.commit()
    return RedirectResponse("/admin/banners", status_code=303)


@router.get("/banners/{banner_id}/products", response_class=HTMLResponse)
async def banner_products_page(
    request: Request,
    banner_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "banners")
    if denied:
        return denied

    banner = (
        await session.execute(
            select(PromoBanner)
            .options(selectinload(PromoBanner.products))
            .where(PromoBanner.id == banner_id)
        )
    ).scalar_one_or_none()
    if not banner:
        return RedirectResponse("/admin/banners", status_code=303)

    attached_ids = {product.id for product in banner.products}
    attached_products = sorted(banner.products, key=lambda product: (product.name_ru or "").lower())
    products_stmt = (
        select(Product)
        .join(Category, Category.id == Product.category_id)
    )
    visible_condition = (Product.is_active == True) & (Category.is_active == True)
    if attached_ids:
        products_stmt = products_stmt.where(
            or_(visible_condition, Product.id.in_(attached_ids))
        )
    else:
        products_stmt = products_stmt.where(visible_condition)
    if attached_ids:
        products_stmt = products_stmt.order_by(
            case((Product.id.in_(attached_ids), 1), else_=0),
            Product.name_ru.asc(),
        )
    else:
        products_stmt = products_stmt.order_by(Product.name_ru.asc())
    products = (await session.execute(products_stmt)).scalars().all()

    return templates.TemplateResponse("admin/banner_products.html", {
        "request": request,
        "user": user,
        "banner": banner,
        "attached_products": attached_products,
        "products": products,
        "attached_ids": attached_ids,
        "csrf_token": generate_csrf_token(request),
    })


@router.post("/banners/{banner_id}/products/{product_id}/toggle")
async def banner_product_toggle(
    request: Request,
    banner_id: int,
    product_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "banners"):
        return RedirectResponse("/admin")

    banner = (
        await session.execute(
            select(PromoBanner)
            .options(selectinload(PromoBanner.products))
            .where(PromoBanner.id == banner_id)
        )
    ).scalar_one_or_none()
    product = (await session.execute(select(Product).where(Product.id == product_id))).scalar_one_or_none()
    if not banner or not product:
        return RedirectResponse("/admin/banners", status_code=303)

    is_attached = any(attached.id == product_id for attached in banner.products)
    if is_attached:
        banner.products = [attached for attached in banner.products if attached.id != product_id]
    else:
        active_category = (
            await session.execute(
                select(Category).where(Category.id == product.category_id, Category.is_active == True)
            )
        ).scalar_one_or_none()
        if not product.is_active or not active_category:
            return RedirectResponse(f"/admin/banners/{banner_id}/products", status_code=303)
        banner.products.append(product)
    await session.commit()

    referer = request.headers.get("referer") or ""
    if f"/admin/banners/{banner_id}/products" in referer:
        return RedirectResponse(referer, status_code=303)
    return RedirectResponse(f"/admin/banners/{banner_id}/products", status_code=303)

@router.post("/banners/{banner_id}/toggle")
async def banner_toggle(
    banner_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "banners"): return RedirectResponse("/admin")
    banner = (await session.execute(select(PromoBanner).where(PromoBanner.id == banner_id))).scalar_one_or_none()
    if banner:
        banner.is_active = not banner.is_active
        await session.commit()
    return RedirectResponse("/admin/banners", status_code=303)

@router.post("/banners/{banner_id}/delete")
async def banner_delete(
    banner_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "banners"): return RedirectResponse("/admin")
    banner = (await session.execute(select(PromoBanner).where(PromoBanner.id == banner_id))).scalar_one_or_none()
    if banner:
        if banner.image_path:
            await delete_file(banner.image_path)
        await session.delete(banner)
        await session.commit()
    return RedirectResponse("/admin/banners", status_code=303)


# ===================== ЗАГРУЗКА ДОП. ФОТО ТОВАРОВ =====================
@router.post("/products/{product_id}/images")
async def product_add_images(
    request: Request,
    product_id: int,
    images: List[UploadFile] = File(...),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "products"): return RedirectResponse("/admin")

    product = await session.get(Product, product_id)
    if not product:
        return RedirectResponse("/admin/products", status_code=303)

    max_sort = (await session.execute(
        select(func.coalesce(func.max(ProductImage.sort_order), 0))
        .where(ProductImage.product_id == product_id)
    )).scalar() or 0

    for idx, img_file in enumerate(images):
        if not img_file.filename:
            continue
        try:
            file_bytes = await read_upload_limited(img_file, IMAGE_MAX_BYTES)
            processed = await asyncio.to_thread(process_product_image, file_bytes)
        except Exception:
            continue

        unique_name = f"{uuid.uuid4()}.{IMAGE_EXTENSION}"
        upload_dir = "media/products"
        await asyncio.to_thread(os.makedirs, upload_dir, exist_ok=True)
        async with aiofiles.open(f"{upload_dir}/{unique_name}", "wb") as f:
            await f.write(processed)

        session.add(ProductImage(
            product_id=product_id,
            image_path=f"/media/products/{unique_name}",
            sort_order=max_sort + idx + 1,
        ))

    await session.commit()
    return RedirectResponse(f"/admin/products/{product_id}/edit", status_code=303)

@router.post("/products/{product_id}/image/delete")
async def product_delete_main_image(
    product_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "products"): return RedirectResponse("/admin")

    product = await session.get(Product, product_id)
    if product and product.image_path:
        old_path = product.image_path
        # Promote first extra image to main, if any
        first_extra = (await session.execute(
            select(ProductImage)
            .where(
                ProductImage.product_id == product_id,
                ProductImage.media_type == "image",
            )
            .order_by(ProductImage.sort_order.asc())
            .limit(1)
        )).scalar_one_or_none()
        if first_extra:
            product.image_path = first_extra.image_path
            await session.delete(first_extra)
        else:
            product.image_path = ""
        await session.commit()
        await delete_file(old_path)
    return RedirectResponse(f"/admin/products/{product_id}/edit", status_code=303)

@router.post("/products/images/{image_id}/delete")
async def product_delete_image(
    image_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf)
):
    if not user: return RedirectResponse("/admin/login")
    if not has_permission(user, "products"): return RedirectResponse("/admin")

    img = (await session.execute(select(ProductImage).where(ProductImage.id == image_id))).scalar_one_or_none()
    if img:
        product_id = img.product_id
        await delete_product_media_files(img.image_path, img.poster_path)
        await session.delete(img)
        await session.commit()
        return RedirectResponse(f"/admin/products/{product_id}/edit", status_code=303)
    return RedirectResponse("/admin/products", status_code=303)

# ===================== ПОДДЕРЖКА (ЧАТ) — АДМИН =====================

SUPPORT_UPLOAD_DIR = "media/support"
SUPPORT_MAX_FILE_SIZE = 10 * 1024 * 1024
SUPPORT_MAX_TEXT_LENGTH = 4000

async def _send_support_notification(telegram_id: int, text: str):
    """Send a Telegram notification to the user about a new support message."""
    try:
        from app.config import settings
        webapp_url = f"{settings.WEB_BASE_URL}/shop/support"
        from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton
        kb = InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="💬 Открыть чат", web_app={"url": webapp_url})
        ]])
        await bot.send_message(telegram_id, text, parse_mode="HTML", reply_markup=kb)
    except Exception:
        logger.opt(exception=True).info(f"Failed to send support notification to {telegram_id}")


@router.get("/support/unread-count")
async def admin_support_unread_count(
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return JSONResponse({"count": 0}, status_code=401)
    if not has_permission(user, "support"):
        return JSONResponse({"count": 0}, status_code=403)

    unread_total = (
        await session.execute(
            select(func.coalesce(func.sum(SupportChat.unread_admin), 0))
        )
    ).scalar() or 0
    return JSONResponse({"count": int(unread_total)})

@router.get("/support", response_class=HTMLResponse)
async def admin_support_list(
    request: Request,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "support")
    if denied: return denied

    stmt = (
        select(SupportChat)
        .options(selectinload(SupportChat.user))
        .order_by(SupportChat.last_message_at.desc().nullslast())
    )
    chats = (await session.execute(stmt)).scalars().all()
    csrf_token = generate_csrf_token(request)
    return templates.TemplateResponse("admin/support_list.html", {
        "request": request, "user": user, "chats": chats, "csrf_token": csrf_token,
    })

@router.get("/support/{chat_id}", response_class=HTMLResponse)
async def admin_support_chat(
    request: Request,
    chat_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "support")
    if denied: return denied

    chat = (await session.execute(
        select(SupportChat).options(selectinload(SupportChat.user)).where(SupportChat.id == chat_id)
    )).scalar_one_or_none()
    if not chat:
        return RedirectResponse("/admin/support")

    chat.unread_admin = 0
    await session.commit()

    ADMIN_SUPPORT_PAGE = 50
    msg_stmt = (
        select(SupportMessage)
        .options(selectinload(SupportMessage.sender))
        .where(SupportMessage.chat_id == chat_id)
        .order_by(SupportMessage.created_at.desc())
        .limit(ADMIN_SUPPORT_PAGE + 1)
    )
    rows = list((await session.execute(msg_stmt)).scalars().all())
    has_older = len(rows) > ADMIN_SUPPORT_PAGE
    if has_older:
        rows = rows[:ADMIN_SUPPORT_PAGE]
    messages = list(reversed(rows))

    csrf_token = generate_csrf_token(request)
    return templates.TemplateResponse("admin/support_chat.html", {
        "request": request, "user": user, "chat": chat,
        "messages": messages, "csrf_token": csrf_token,
        "has_older": has_older,
        "oldest_id": messages[0].id if messages else 0,
    })

@router.post("/support/{chat_id}/send")
async def admin_support_send(
    request: Request,
    chat_id: int,
    text: str = Form(""),
    image: UploadFile = File(None),
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return JSONResponse({"error": "unauthorized"}, status_code=401)
    if not has_permission(user, "support"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    chat = (await session.execute(
        select(SupportChat).options(selectinload(SupportChat.user)).where(SupportChat.id == chat_id)
    )).scalar_one_or_none()
    if not chat:
        return JSONResponse({"error": "not found"}, status_code=404)

    text = text.strip()
    if len(text) > SUPPORT_MAX_TEXT_LENGTH:
        return JSONResponse({"error": "Message too long"}, status_code=400)

    image_path = None

    if image and image.filename:
        try:
            content = await read_upload_limited(image, SUPPORT_MAX_FILE_SIZE)
            ext = validate_image_upload(image.filename, content)
        except ValueError as exc:
            if str(exc) == "file_too_large":
                return JSONResponse({"error": "File too large"}, status_code=400)
            return JSONResponse({"error": "Invalid image"}, status_code=400)
        os.makedirs(SUPPORT_UPLOAD_DIR, exist_ok=True)
        fname = f"{uuid.uuid4().hex}{ext}"
        fpath = os.path.join(SUPPORT_UPLOAD_DIR, fname)
        async with aiofiles.open(fpath, "wb") as f:
            await f.write(content)
        image_path = fpath

    if not text and not image_path:
        return JSONResponse({"error": "empty"}, status_code=400)

    msg = SupportMessage(
        chat_id=chat.id,
        sender_type="admin",
        sender_id=user.id,
        text=text or None,
        image_path=image_path,
    )
    session.add(msg)
    chat.last_message_at = datetime.utcnow()
    chat.unread_user += 1
    await session.commit()

    # Notify user via Telegram bot
    if chat.user and chat.user.telegram_id:
        lang = chat.user.language or "ru"
        asyncio.create_task(_send_support_notification(chat.user.telegram_id, tr("bot_support_reply", lang)))

    return JSONResponse({
        "ok": True,
        "message": {
            "id": msg.id,
            "sender_type": "admin",
            "sender_name": user.login or user.username or "Админ",
            "text": msg.text,
            "image_path": f"/{msg.image_path}" if msg.image_path else None,
            "created_at": msg.created_at.strftime("%H:%M"),
        }
    })

@router.get("/support/{chat_id}/messages")
async def admin_support_messages(
    request: Request,
    chat_id: int,
    after: int = 0,
    before: int = 0,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return JSONResponse({"error": "unauthorized"}, status_code=401)
    if not has_permission(user, "support"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    chat = (await session.execute(
        select(SupportChat).where(SupportChat.id == chat_id)
    )).scalar_one_or_none()
    if not chat:
        return {"messages": [], "has_older": False}

    ADMIN_PAGE = 50

    if before > 0:
        msg_stmt = (
            select(SupportMessage)
            .options(selectinload(SupportMessage.sender))
            .where(SupportMessage.chat_id == chat_id, SupportMessage.id < before)
            .order_by(SupportMessage.created_at.desc())
            .limit(ADMIN_PAGE + 1)
        )
        rows = list((await session.execute(msg_stmt)).scalars().all())
        has_older = len(rows) > ADMIN_PAGE
        if has_older:
            rows = rows[:ADMIN_PAGE]
        msgs = list(reversed(rows))
        return {
            "has_older": has_older,
            "messages": [
                {
                    "id": m.id,
                    "sender_type": m.sender_type,
                    "sender_name": (m.sender.login or m.sender.username or "Админ") if m.sender else "Админ",
                    "text": m.text,
                    "image_path": f"/{m.image_path}" if m.image_path else None,
                    "created_at": m.created_at.strftime("%H:%M"),
                }
                for m in msgs
            ]
        }

    if after > 0:
        chat.unread_admin = 0
        await session.commit()

    msg_stmt = (
        select(SupportMessage)
        .options(selectinload(SupportMessage.sender))
        .where(SupportMessage.chat_id == chat_id, SupportMessage.id > after)
        .order_by(SupportMessage.created_at.asc())
    )
    msgs = (await session.execute(msg_stmt)).scalars().all()
    return {
        "messages": [
            {
                "id": m.id,
                "sender_type": m.sender_type,
                "sender_name": (m.sender.login or m.sender.username or "Админ") if m.sender else "Админ",
                "text": m.text,
                "image_path": f"/{m.image_path}" if m.image_path else None,
                "created_at": m.created_at.strftime("%H:%M"),
            }
            for m in msgs
        ]
    }

@router.post("/support/{chat_id}/close")
async def admin_support_close(
    chat_id: int,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "support"): return RedirectResponse("/admin")
    chat = (await session.execute(
        select(SupportChat).where(SupportChat.id == chat_id)
    )).scalar_one_or_none()
    if chat:
        chat.is_closed = True
        await session.commit()
    return RedirectResponse("/admin/support", status_code=303)


# ─────────────────────── АНАЛИТИКА ───────────────────────

@router.get("/analytics", response_class=HTMLResponse)
async def analytics_page(
    request: Request,
    month: Optional[int] = None,
    year: Optional[int] = None,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "analytics")
    if denied: return denied

    uz_now = datetime.now(timezone(timedelta(hours=5)))
    sel_year = year or uz_now.year
    sel_month = month or uz_now.month

    # UTC boundaries of selected month (Tashkent UTC+5)
    month_start_local = date(sel_year, sel_month, 1)
    if sel_month == 12:
        month_end_local = date(sel_year + 1, 1, 1)
    else:
        month_end_local = date(sel_year, sel_month + 1, 1)
    month_start_utc = datetime(month_start_local.year, month_start_local.month, month_start_local.day) - timedelta(hours=5)
    month_end_utc = datetime(month_end_local.year, month_end_local.month, month_end_local.day) - timedelta(hours=5)

    # ── 1. Revenue & AOV ──
    rev_stmt = select(
        func.coalesce(func.sum(Order.total_amount), 0),
        func.count(Order.id),
    ).where(
        Order.status.in_(["done", "paid"]),
        Order.created_at >= month_start_utc,
        Order.created_at < month_end_utc,
    )
    rev_row = (await session.execute(rev_stmt)).one()
    total_revenue = int(rev_row[0])
    completed_orders = int(rev_row[1])
    aov = total_revenue // completed_orders if completed_orders else 0

    # ── 2. Order conversion ──
    conv_stmt = select(Order.status, func.count(Order.id)).where(
        Order.created_at >= month_start_utc,
        Order.created_at < month_end_utc,
    ).group_by(Order.status)
    conv_raw = (await session.execute(conv_stmt)).all()
    conv_map = {r[0]: r[1] for r in conv_raw}
    orders_done = conv_map.get("done", 0) + conv_map.get("paid", 0)
    orders_cancelled = conv_map.get("cancelled", 0)
    orders_new = conv_map.get("new", 0)
    orders_delivery = conv_map.get("delivery", 0)
    orders_total = sum(conv_map.values())

    # ── 3. Daily sales ──
    days_in_month = calendar.monthrange(sel_year, sel_month)[1]
    daily_stmt = select(
        func.extract("day", Order.created_at + timedelta(hours=5)).label("d"),
        func.sum(Order.total_amount),
        func.count(Order.id),
    ).where(
        Order.status.in_(["done", "paid"]),
        Order.created_at >= month_start_utc,
        Order.created_at < month_end_utc,
    ).group_by("d").order_by("d")
    daily_raw = (await session.execute(daily_stmt)).all()
    daily_map = {int(r[0]): (int(r[1]), int(r[2])) for r in daily_raw}
    daily_labels = list(range(1, days_in_month + 1))
    daily_revenue = [daily_map.get(d, (0, 0))[0] for d in daily_labels]
    daily_count = [daily_map.get(d, (0, 0))[1] for d in daily_labels]

    # ── 4. Top products (ABC-light) ──
    top_stmt = (
        select(
            OrderItem.product_name,
            func.sum(OrderItem.quantity).label("qty"),
            func.sum(OrderItem.price_at_purchase * OrderItem.quantity).label("rev"),
        )
        .join(Order, OrderItem.order_id == Order.id)
        .where(
            Order.status.in_(["done", "paid"]),
            Order.created_at >= month_start_utc,
            Order.created_at < month_end_utc,
        )
        .group_by(OrderItem.product_name)
        .order_by(func.sum(OrderItem.price_at_purchase * OrderItem.quantity).desc())
        .limit(10)
    )
    top_raw = (await session.execute(top_stmt)).all()
    top_products = [
        {"name": r[0], "qty": int(r[1]), "revenue": int(r[2])}
        for r in top_raw
    ]
    top_prod_labels = [p["name"] for p in top_products[:7]]
    top_prod_rev = [p["revenue"] for p in top_products[:7]]
    top_prod_qty = [p["qty"] for p in top_products[:7]]

    # ── 5. Payment methods ──
    pay_stmt = select(
        Order.payment_method, func.count(Order.id), func.coalesce(func.sum(Order.total_amount), 0)
    ).where(
        Order.status.in_(["done", "paid"]),
        Order.created_at >= month_start_utc,
        Order.created_at < month_end_utc,
    ).group_by(Order.payment_method)
    pay_raw = (await session.execute(pay_stmt)).all()
    pay_name_map = {"cash": "Наличные", "card": "Payme", "click": "Click", "card_transfer": "Перевод на карту"}
    payment_stats = [
        {"method": pay_name_map.get(r[0], r[0]), "count": int(r[1]), "amount": int(r[2])}
        for r in pay_raw
    ]
    pay_labels = [p["method"] for p in payment_stats]
    pay_data = [p["count"] for p in payment_stats]
    pay_amount = [p["amount"] for p in payment_stats]

    # ── 6. New customers ──
    new_cust_stmt = select(func.count(User.id)).where(
        User.role == "user",
        User.created_at >= month_start_utc,
        User.created_at < month_end_utc,
    )
    new_customers = (await session.execute(new_cust_stmt)).scalar() or 0

    # Month selector helpers
    months_map = {
        1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
        5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
        9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь",
    }

    return templates.TemplateResponse("admin/analytics.html", {
        "request": request,
        "user": user,
        "csrf_token": generate_csrf_token(request),
        # Period
        "sel_year": sel_year,
        "sel_month": sel_month,
        "sel_month_name": months_map.get(sel_month, ""),
        "months_map": months_map,
        "current_year": uz_now.year,
        # 1 Revenue
        "total_revenue": total_revenue,
        "aov": aov,
        "completed_orders": completed_orders,
        # 2 Conversion
        "orders_total": orders_total,
        "orders_done": orders_done,
        "orders_cancelled": orders_cancelled,
        "orders_new": orders_new,
        "orders_delivery": orders_delivery,
        # 3 Daily
        "daily_labels": daily_labels,
        "daily_revenue": daily_revenue,
        "daily_count": daily_count,
        # 4 Top products
        "top_products": top_products,
        "top_prod_labels": top_prod_labels,
        "top_prod_rev": top_prod_rev,
        "top_prod_qty": top_prod_qty,
        # 5 Payment
        "payment_stats": payment_stats,
        "pay_labels": pay_labels,
        "pay_data": pay_data,
        "pay_amount": pay_amount,
        # 6 New customers
        "new_customers": new_customers,
    })


@router.get("/analytics/export")
async def analytics_export(
    request: Request,
    month: Optional[int] = None,
    year: Optional[int] = None,
    user: User = Depends(get_current_admin),
    session: AsyncSession = Depends(get_db),
):
    """Export monthly analytics as XLSX."""
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "analytics"): return RedirectResponse("/admin")

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from fastapi.responses import StreamingResponse

    uz_now = datetime.now(timezone(timedelta(hours=5)))
    sel_year = year or uz_now.year
    sel_month = month or uz_now.month

    month_start_local = date(sel_year, sel_month, 1)
    if sel_month == 12:
        month_end_local = date(sel_year + 1, 1, 1)
    else:
        month_end_local = date(sel_year, sel_month + 1, 1)
    month_start_utc = datetime(month_start_local.year, month_start_local.month, month_start_local.day) - timedelta(hours=5)
    month_end_utc = datetime(month_end_local.year, month_end_local.month, month_end_local.day) - timedelta(hours=5)

    months_map = {
        1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
        5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
        9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь",
    }

    # Fetch all orders in period
    orders_stmt = (
        select(Order)
        .options(selectinload(Order.items), selectinload(Order.user))
        .where(
            Order.created_at >= month_start_utc,
            Order.created_at < month_end_utc,
        )
        .order_by(Order.created_at)
    )
    orders = (await session.execute(orders_stmt)).scalars().all()

    # New customers
    new_cust_stmt = select(func.count(User.id)).where(
        User.role == "user",
        User.created_at >= month_start_utc,
        User.created_at < month_end_utc,
    )
    new_customers = (await session.execute(new_cust_stmt)).scalar() or 0

    wb = Workbook()
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="3E2310", end_color="3E2310", fill_type="solid")
    subheader_font = Font(bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    pay_name_map = {"cash": "Наличные", "card": "Payme", "click": "Click", "card_transfer": "Перевод на карту"}
    status_name_map = {"new": "Новый", "paid": "Оплачен", "delivery": "В пути", "done": "Завершён", "cancelled": "Отменён"}

    # ── Sheet 1: Summary ──
    ws = wb.active
    ws.title = "Сводка"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25
    title = f"Аналитика — {months_map.get(sel_month, '')} {sel_year}"
    ws.append([title])
    ws.merge_cells("A1:B1")
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    summary_data = [
        ("Общая выручка (сум)", sum(o.total_amount for o in orders if o.status in ("done", "paid"))),
        ("Кол-во завершённых заказов", sum(1 for o in orders if o.status in ("done", "paid"))),
        ("Средний чек (сум)", sum(o.total_amount for o in orders if o.status in ("done", "paid")) // max(1, sum(1 for o in orders if o.status in ("done", "paid")))),
        ("Всего заказов", len(orders)),
        ("Отменённых заказов", sum(1 for o in orders if o.status == "cancelled")),
        ("Новые клиенты", new_customers),
    ]
    for label, val in summary_data:
        ws.append([label, val])

    # ── Sheet 2: Orders list ──
    ws2 = wb.create_sheet("Заказы")
    cols = ["ID", "Дата", "Клиент", "Телефон", "Статус", "Способ оплаты", "Доставка", "Стоимость доставки", "Адрес", "Сумма (сум)", "Скидка (сум)", "Комментарий"]
    ws2.append(cols)
    for i, cell in enumerate(ws2[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for o in orders:
        local_dt = o.created_at + timedelta(hours=5)
        ws2.append([
            o.id,
            local_dt.strftime("%d.%m.%Y %H:%M"),
            o.user.username if o.user else (o.customer_name or ""),
            o.contact_phone or "",
            status_name_map.get(o.status, o.status),
            pay_name_map.get(o.payment_method, o.payment_method),
            "Самовывоз" if o.delivery_method == "pickup" else (o.delivery_option_name or "Доставка"),
            o.delivery_price or 0,
            o.delivery_address or "",
            o.total_amount,
            o.discount_amount,
            o.comment or "",
        ])

    for col_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]:
        ws2.column_dimensions[col_letter].width = 18

    # ── Sheet 3: Top products ──
    ws3 = wb.create_sheet("Топ товаров")
    ws3.append(["Товар", "Кол-во (шт)", "Выручка (сум)"])
    for i, cell in enumerate(ws3[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Aggregate from order items of completed orders
    top_agg: dict = {}
    for o in orders:
        if o.status not in ("done", "paid"):
            continue
        for item in o.items:
            key = item.product_name
            if key not in top_agg:
                top_agg[key] = {"qty": 0, "rev": 0}
            top_agg[key]["qty"] += item.quantity
            top_agg[key]["rev"] += item.price_at_purchase * item.quantity
    for name, data in sorted(top_agg.items(), key=lambda x: x[1]["rev"], reverse=True):
        ws3.append([name, data["qty"], data["rev"]])
    ws3.column_dimensions["A"].width = 35
    ws3.column_dimensions["B"].width = 15
    ws3.column_dimensions["C"].width = 20

    # ── Sheet 4: Payment methods ──
    ws4 = wb.create_sheet("Способы оплаты")
    ws4.append(["Способ оплаты", "Кол-во заказов", "Сумма (сум)"])
    for i, cell in enumerate(ws4[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
    pay_agg: dict = {}
    for o in orders:
        if o.status not in ("done", "paid"):
            continue
        pm = pay_name_map.get(o.payment_method, o.payment_method)
        if pm not in pay_agg:
            pay_agg[pm] = {"count": 0, "amount": 0}
        pay_agg[pm]["count"] += 1
        pay_agg[pm]["amount"] += o.total_amount
    for pm, data in pay_agg.items():
        ws4.append([pm, data["count"], data["amount"]])
    ws4.column_dimensions["A"].width = 20
    ws4.column_dimensions["B"].width = 20
    ws4.column_dimensions["C"].width = 20

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"analytics_{sel_year}_{sel_month:02d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ===================== БЭКАП БАЗЫ ДАННЫХ =====================

@router.get("/backup", response_class=HTMLResponse)
@router.get("/backup/", response_class=HTMLResponse)
async def backup_page(
    request: Request,
    user: User = Depends(get_current_admin),
):
    if not user:
        return RedirectResponse("/admin/login")
    denied = check_permission(request, user, "backup")
    if denied: return denied

    return templates.TemplateResponse("admin/backup.html", {
        "request": request,
        "user": user,
        "rollback_state": backup_service.get_rollback_state(),
        "csrf_token": generate_csrf_token(request),
    })


@router.post("/backup/export/start")
async def backup_export_start(
    user: User = Depends(get_current_admin),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return JSONResponse({"error": "forbidden"}, status_code=403)
    if not has_permission(user, "backup"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    try:
        job = await backup_service.start_export_job()
    except BackupBusyError as exc:
        return JSONResponse({"error": str(exc)}, status_code=409)
    except Exception as exc:
        logger.exception("Backup export start error")
        return JSONResponse({"error": str(exc)}, status_code=500)

    return JSONResponse(job, status_code=202)


@router.post("/backup/import/start")
async def backup_import(
    dump_file: UploadFile = File(...),
    user: User = Depends(get_current_admin),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return JSONResponse({"error": "forbidden"}, status_code=403)
    if not has_permission(user, "backup"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    if not dump_file or not dump_file.filename:
        return JSONResponse({"error": "Файл не выбран"}, status_code=400)

    fname = dump_file.filename.lower()
    if not fname.endswith(".json") and not fname.endswith(".zip"):
        return JSONResponse(
            {"error": "Поддерживаются только .zip и legacy .json файлы"},
            status_code=400,
        )

    upload_path = BACKUP_UPLOAD_DIR / f"{uuid.uuid4().hex}{Path(dump_file.filename).suffix.lower()}"
    try:
        if backup_service.is_busy():
            return JSONResponse({"error": "Another backup job is already running"}, status_code=409)
        await backup_service.stream_upload_to_disk(dump_file, upload_path)
        job = await backup_service.start_import_job(upload_path, dump_file.filename)
        return JSONResponse(job, status_code=202)
    except BackupBusyError as exc:
        await asyncio.to_thread(upload_path.unlink, missing_ok=True)
        return JSONResponse({"error": str(exc)}, status_code=409)
    except ValueError as exc:
        await asyncio.to_thread(upload_path.unlink, missing_ok=True)
        return JSONResponse({"error": str(exc)}, status_code=400)
    except Exception as exc:
        logger.exception("Backup import start error")
        await asyncio.to_thread(upload_path.unlink, missing_ok=True)
        return JSONResponse({"error": str(exc)}, status_code=500)


@router.post("/backup/restore/start")
async def backup_restore_previous(
    user: User = Depends(get_current_admin),
    csrf: bool = Depends(validate_csrf),
):
    if not user:
        return JSONResponse({"error": "forbidden"}, status_code=403)
    if not has_permission(user, "backup"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    try:
        job = await backup_service.start_restore_previous_job()
        return JSONResponse(job, status_code=202)
    except FileNotFoundError as exc:
        return JSONResponse({"error": str(exc)}, status_code=404)
    except BackupBusyError as exc:
        return JSONResponse({"error": str(exc)}, status_code=409)
    except Exception as exc:
        logger.exception("Backup restore start error")
        return JSONResponse({"error": str(exc)}, status_code=500)


@router.get("/backup/jobs/{job_id}")
async def backup_job_status(
    job_id: str,
    user: User = Depends(get_current_admin)
):
    if not user:
        return JSONResponse({"error": "forbidden"}, status_code=403)
    if not has_permission(user, "backup"):
        return JSONResponse({"error": "forbidden"}, status_code=403)

    job = backup_service.get_job(job_id)
    if not job:
        return JSONResponse({"error": "not_found"}, status_code=404)
    return JSONResponse(job)


@router.get("/backup/jobs/{job_id}/download")
async def backup_download(
    request: Request,
    job_id: str,
    user: User = Depends(get_current_admin),
):
    if not user:
        return RedirectResponse("/admin/login")
    if not has_permission(user, "backup"):
        return RedirectResponse("/admin")

    job = backup_service.get_job(job_id)
    if not job:
        return JSONResponse({"error": "not_found"}, status_code=404)
    if job.get("kind") != "export":
        return JSONResponse({"error": "invalid_job_kind"}, status_code=400)
    if job.get("status") != "done":
        return JSONResponse({"error": "backup_not_ready"}, status_code=409)

    ready_file = backup_service.resolve_ready_file(job_id)
    if not ready_file:
        return JSONResponse({"error": "file_not_found"}, status_code=404)

    filename = job.get("filename") or ready_file.name
    content_disposition = f'attachment; filename="{filename}"'

    if request.headers.get("x-forwarded-proto"):
        return Response(
            status_code=200,
            headers={
                "X-Accel-Redirect": f"/_protected_backups/{ready_file.name}",
                "Content-Disposition": content_disposition,
                "Content-Type": "application/zip",
                "Cache-Control": "no-store",
            },
        )

    return FileResponse(
        path=ready_file,
        media_type="application/zip",
        filename=filename,
        headers={"Cache-Control": "no-store"},
    )
